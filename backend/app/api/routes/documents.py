import uuid
from pathlib import Path
from typing import Any
import json
import time
import base64
import tempfile
import zipfile
import rarfile
from io import BytesIO
from xml.sax.saxutils import escape
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import FileResponse
from sqlmodel import col, select, Session
from sqlalchemy import or_, func

from app.api.deps import CurrentUser, SessionDep
from datetime import datetime, timezone
from app.models import (
    Document,
    DocumentChunk,
    DocumentChunkPublic,
    DocumentChunksPublic,
    DocumentPublic,
    DocumentsPublic,
    KnowledgeBase,
    Message,
    KnowledgeBaseSearchRequest,
    KnowledgeBaseSearchResult,
    KnowledgeBaseSearchResults,
    RagChatRequest,
    RagChatResponse,
    RagChatSource,
    KnowledgeBaseSemanticSearchRequest,
    KnowledgeBaseSemanticSearchResult,
    KnowledgeBaseSemanticSearchResults,
    KnowledgeBaseEmbeddingBackfillRequest,
    KnowledgeBaseEmbeddingBackfillResult,
    RagRun,
    RagRunPublic,
    RagRunsPublic,
    RagEvalCase,
    RagEvalCaseCreate,
    RagEvalCasePublic,
    RagEvalCasesPublic,
    RagEvalRun,
    RagEvalRunPublic,
    RagEvalRunRequest,
    RagEvalRunsPublic,
    RagEvalBatchRunRequest,
    RagEvalBatchRunResult,
    RagEvalSummary,
    RagEvalFailureAnalysis,
    RagEvalParamGroup,
    RagEvalParamGroups,
    RagRetrievalPreset,
    RagRetrievalPresetCreate,
    RagRetrievalPresetPublic,
    RagRetrievalPresetsPublic,
    RagEvalBatch,
    RagEvalBatchPublic,
    RagEvalBatchesPublic,
    AgentChatRequest,
    AgentChatResponse,
    AgentRun,
    AgentRunPublic,
    AgentRunsPublic,
    AgentToolCallPublic,
    RagAgentCompareEvalRequest,
    RagAgentCompareEvalItem,
    RagAgentCompareEvalSummary,
    RagAgentCompareEvalResponse,
    RagAgentCompareBatch,
    RagAgentCompareItem,
    RagAgentCompareBatchPublic,
    RagAgentCompareBatchesPublic,
    RagAgentCompareItemPublic,
    RagAgentCompareBatchDetailPublic,
    KnowledgeBaseAgentSettings,
    KnowledgeBaseAgentSettingsPublic,
    KnowledgeBaseAgentSettingsUpdate,
    RagAgentCompareReportResponse,
    RagAgentCompareFileReportResponse,
    CodeEvalTypeCompareStat,
    CodeEvalTypeCompareSummaryPublic,
    RagAgentCompareFailureReasonStat,
    RagAgentCompareFailureCasePublic,
    RagAgentCompareFailureAnalysisPublic,
    CodeAgentExperimentReportPublic,
    RepositoryAnalysisTask,
)
from app.services.rag import (
    build_search_terms,
    count_keyword_matches,
    generate_rag_answer,
    retrieve_chunks_for_rag,
    hybrid_retrieve_chunks_for_rag,
    build_rag_chat_source,
)
from app.services.embedding import (
    EmbeddingError,
    backfill_chunk_embeddings,
    embed_document_chunk,
    semantic_search_chunks,
)
from app.services.agent import run_knowledge_base_agent
from app.services.document_parser import (
    DocumentParseError,
    parse_document_content,
    SUPPORTED_DOCUMENT_EXTENSIONS,
)
from app.services.code_parser import is_code_file, split_code_into_chunks, CODE_EXTENSIONS

router = APIRouter(tags=["documents"])

STORAGE_ROOT = Path("storage/documents")
MAX_FILE_SIZE = 1024 * 1024 * 1024
MAX_REPOSITORY_ARCHIVE_SIZE = 1024 * 1024 * 1024
MAX_REPOSITORY_FILE_COUNT = 3000
MAX_REPOSITORY_TOTAL_CODE_SIZE = 1024 * 1024 * 1024
UPLOAD_COPY_CHUNK_SIZE = 1024 * 1024
MAX_REPOSITORY_SINGLE_CODE_FILE_SIZE = 2 * 1024 * 1024

ALLOWED_EXTENSIONS = (
    set(SUPPORTED_DOCUMENT_EXTENSIONS)
    | set(CODE_EXTENSIONS.keys())
)

ALLOWED_CONTENT_TYPES = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".txt": {"text/plain", "application/octet-stream"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
}

TEXT_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
REPOSITORY_IGNORE_DIRS = {
    ".git",
    ".idea",
    ".vscode",
    "__pycache__",
    "node_modules",
    "dist",
    "build",
    "out",
    "target",
    ".next",
    ".nuxt",
    ".venv",
    "venv",
    "env",
    ".pytest_cache",
    ".mypy_cache",
    ".ruff_cache",
    "coverage",
}

REPOSITORY_IGNORE_FILE_PREFIXES = {
    ".DS_Store",
}

REPOSITORY_IGNORE_FILE_SUFFIXES = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".ico",
    ".svg",
    ".mp4",
    ".mp3",
    ".wav",
    ".pdf",
    ".docx",
    ".zip",
    ".rar",
    ".7z",
    ".exe",
    ".dll",
    ".so",
    ".dylib",
    ".class",
    ".jar",
    ".lock",
}

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200


def get_datetime_utc() -> datetime:
    return datetime.now(timezone.utc)


def get_knowledge_base_or_404(
    *,
    session: SessionDep,
    knowledge_base_id: uuid.UUID,
) -> KnowledgeBase:
    knowledge_base = session.get(KnowledgeBase, knowledge_base_id)
    if not knowledge_base:
        raise HTTPException(status_code=404, detail="Knowledge base not found")
    return knowledge_base


def check_knowledge_base_permission(
    *,
    knowledge_base: KnowledgeBase,
    current_user: CurrentUser,
) -> None:
    if not current_user.is_superuser and knowledge_base.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not enough permissions")


def resolve_repository_analysis_scope(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    repository_analysis_task_id:
        uuid.UUID | None,
) -> uuid.UUID | None:
    if (
        repository_analysis_task_id
        is None
    ):
        return None

    task = session.get(
        RepositoryAnalysisTask,
        repository_analysis_task_id,
    )

    if task is None:
        raise HTTPException(
            status_code=404,
            detail={
                "code":
                    "REPOSITORY_ANALYSIS_NOT_FOUND",
                "message":
                    "Repository analysis task not found",
            },
        )

    if (
        not current_user.is_superuser
        and task.owner_id
        != current_user.id
    ):
        raise HTTPException(
            status_code=403,
            detail={
                "code":
                    "REPOSITORY_ANALYSIS_FORBIDDEN",
                "message":
                    "You do not have permission to use this repository analysis task",
            },
        )

    if (
        task.knowledge_base_id
        != knowledge_base_id
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code":
                    "REPOSITORY_ANALYSIS_KNOWLEDGE_BASE_MISMATCH",
                "message":
                    "Repository analysis task does not belong to this knowledge base",
            },
        )

    if (
        task.status != "completed"
        or task.analysis_mode
        != "deep"
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code":
                    "REPOSITORY_ANALYSIS_NOT_READY_FOR_CHAT",
                "message":
                    "Repository analysis task is not ready for repository-scoped chat",
            },
        )

    return task.id

def get_document_or_404(
    *,
    session: SessionDep,
    document_id: uuid.UUID,
) -> Document:
    document = session.get(Document, document_id)
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    return document


def check_document_permission(
    *,
    document: Document,
    current_user: CurrentUser,
) -> None:
    if not current_user.is_superuser and document.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not enough permissions")


def rag_run_to_public(rag_run: RagRun) -> RagRunPublic:
    try:
        source_items = json.loads(rag_run.sources_json or "[]")
    except json.JSONDecodeError:
        source_items = []

    sources = []

    for item in source_items:
        try:
            sources.append(RagChatSource.model_validate(item))
        except Exception:
            continue

    try:
        trace = json.loads(rag_run.trace_json or "[]")
    except json.JSONDecodeError:
        trace = []

    if not isinstance(trace, list):
        trace = []

    return RagRunPublic(
        id=rag_run.id,
        knowledge_base_id=rag_run.knowledge_base_id,
        owner_id=rag_run.owner_id,
        repository_analysis_task_id=rag_run.repository_analysis_task_id,
        source_commit_sha=rag_run.source_commit_sha,
        question=rag_run.question,
        answer=rag_run.answer,
        retrieval_type=rag_run.retrieval_type,
        top_k=rag_run.top_k,
        semantic_weight=rag_run.semantic_weight,
        keyword_weight=rag_run.keyword_weight,
        latency_ms=rag_run.latency_ms,
        sources=sources,
        trace=trace,
        error_message=rag_run.error_message,
        created_at=rag_run.created_at,
    )


def dump_json_data(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def dump_model_list(items: list[Any]) -> str:
    data: list[Any] = []

    for item in items:
        if hasattr(item, "model_dump"):
            data.append(item.model_dump(mode="json"))
        else:
            data.append(item)

    return dump_json_data(data)


def parse_json_data(raw: str | None, default: Any) -> Any:
    if not raw:
        return default

    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return default


def parse_rag_sources_from_json(raw: str | None) -> list[RagChatSource]:
    data = parse_json_data(raw, [])

    if not isinstance(data, list):
        return []

    sources: list[RagChatSource] = []

    for item in data:
        if not isinstance(item, dict):
            continue

        try:
            sources.append(RagChatSource(**item))
        except Exception:
            continue

    return sources


def parse_agent_tool_calls_from_json(
    raw: str | None,
) -> list[AgentToolCallPublic]:
    data = parse_json_data(raw, [])

    if not isinstance(data, list):
        return []

    tool_calls: list[AgentToolCallPublic] = []

    for item in data:
        if not isinstance(item, dict):
            continue

        try:
            tool_calls.append(AgentToolCallPublic(**item))
        except Exception:
            continue

    return tool_calls


def build_compare_batch_public(
    batch: RagAgentCompareBatch,
) -> RagAgentCompareBatchPublic:
    return RagAgentCompareBatchPublic(
        id=batch.id,
        knowledge_base_id=batch.knowledge_base_id,
        owner_id=batch.owner_id,
        name=batch.name,
        top_k=batch.top_k,
        max_steps=batch.max_steps,
        limit=batch.limit,
        semantic_weight=batch.semantic_weight,
        keyword_weight=batch.keyword_weight,
        total_cases=batch.total_cases,
        ran=batch.ran,
        failed=batch.failed,
        average_rag_latency_ms=batch.average_rag_latency_ms,
        average_agent_latency_ms=batch.average_agent_latency_ms,
        average_rag_sources_count=batch.average_rag_sources_count,
        average_agent_sources_count=batch.average_agent_sources_count,
        average_agent_tool_call_count=batch.average_agent_tool_call_count,
        total_agent_failed_tool_count=batch.total_agent_failed_tool_count,
        summarize_document_count=batch.summarize_document_count,
        read_document_chunks_count=batch.read_document_chunks_count,
        search_rag_history_count=batch.search_rag_history_count,
        created_at=batch.created_at,
    )


def build_compare_item_public(
    item: RagAgentCompareItem,
) -> RagAgentCompareItemPublic:
    return RagAgentCompareItemPublic(
        id=item.id,
        batch_id=item.batch_id,
        knowledge_base_id=item.knowledge_base_id,
        owner_id=item.owner_id,
        eval_case_id=item.eval_case_id,
        question=item.question,
        expected_keywords=parse_json_data(
            item.expected_keywords_json,
            [],
        ),
        expected_source_filename=item.expected_source_filename,
        rag_answer=item.rag_answer,
        agent_answer=item.agent_answer,
        rag_latency_ms=item.rag_latency_ms,
        agent_latency_ms=item.agent_latency_ms,
        rag_sources_count=item.rag_sources_count,
        agent_sources_count=item.agent_sources_count,
        agent_tool_call_count=item.agent_tool_call_count,
        agent_failed_tool_count=item.agent_failed_tool_count,
        agent_tool_names=parse_json_data(
            item.agent_tool_names_json,
            [],
        ),
        rag_sources=parse_rag_sources_from_json(item.rag_sources_json),
        agent_sources=parse_rag_sources_from_json(item.agent_sources_json),
        agent_tool_calls=parse_agent_tool_calls_from_json(
            item.agent_tool_calls_json,
        ),
        rag_trace=parse_json_data(item.rag_trace_json, []),
        agent_trace=parse_json_data(item.agent_trace_json, []),
        rag_error_message=item.rag_error_message,
        agent_error_message=item.agent_error_message,
        is_failed=item.is_failed,
        created_at=item.created_at,
    )


def truncate_report_text(text: str | None, max_length: int = 800) -> str:
    if not text:
        return "无"

    text = text.strip()

    if len(text) <= max_length:
        return text

    return text[:max_length] + "\n\n……（内容过长，已截断）"


def format_report_number(value: float | int | None, suffix: str = "") -> str:
    if value is None:
        return "-"

    if isinstance(value, float):
        return f"{value:.2f}{suffix}"

    return f"{value}{suffix}"


def build_agent_tool_chain_text(item: RagAgentCompareItemPublic) -> str:
    if not item.agent_tool_names:
        return "无工具调用"

    return " → ".join(item.agent_tool_names)


def build_item_warning_reasons(item: RagAgentCompareItemPublic) -> list[str]:
    reasons: list[str] = []

    if item.is_failed:
        reasons.append("整体标记失败")

    if item.rag_error_message:
        reasons.append("RAG 运行出错")

    if item.agent_error_message:
        reasons.append("Agent 运行出错")

    if item.agent_failed_tool_count > 0:
        reasons.append(f"Agent 工具失败 {item.agent_failed_tool_count} 次")

    if item.rag_sources_count == 0:
        reasons.append("RAG 无 sources")

    if item.agent_sources_count == 0:
        reasons.append("Agent 无 sources")

    if item.agent_tool_call_count == 0:
        reasons.append("Agent 无工具调用")

    if (
        item.rag_latency_ms is not None
        and item.agent_latency_ms is not None
        and (
            item.agent_latency_ms > item.rag_latency_ms * 2
            or item.agent_latency_ms - item.rag_latency_ms > 3000
        )
    ):
        reasons.append("Agent 明显慢于 RAG")

    return reasons

def build_base64_file_response(
    *,
    filename: str,
    mime_type: str,
    file_bytes: bytes,
) -> RagAgentCompareFileReportResponse:
    return RagAgentCompareFileReportResponse(
        filename=filename,
        mime_type=mime_type,
        content_base64=base64.b64encode(file_bytes).decode("utf-8"),
    )


def build_rag_agent_compare_markdown_report(
    *,
    knowledge_base_name: str,
    batch: RagAgentCompareBatchPublic,
    items: list[RagAgentCompareItemPublic],
) -> str:
    failed_items = [
        item for item in items
        if item.is_failed
        or item.rag_error_message
        or item.agent_error_message
        or item.agent_failed_tool_count > 0
        or item.rag_sources_count == 0
        or item.agent_sources_count == 0
    ]

    slow_agent_items = [
        item for item in items
        if item.rag_latency_ms is not None
        and item.agent_latency_ms is not None
        and (
            item.agent_latency_ms > item.rag_latency_ms * 2
            or item.agent_latency_ms - item.rag_latency_ms > 3000
        )
    ]

    lines: list[str] = []

    lines.append("# RAG vs Agent 实验报告")
    lines.append("")
    lines.append("## 一、实验基本信息")
    lines.append("")
    lines.append(f"- 知识库：{knowledge_base_name}")
    lines.append(f"- 批次名称：{batch.name}")
    lines.append(f"- 批次 ID：{batch.id}")
    lines.append(
        f"- 创建时间：{batch.created_at.isoformat() if batch.created_at else '-'}"
    )
    lines.append(f"- top_k：{batch.top_k}")
    lines.append(f"- max_steps：{batch.max_steps}")
    lines.append(f"- semantic_weight：{batch.semantic_weight}")
    lines.append(f"- keyword_weight：{batch.keyword_weight}")
    lines.append(f"- limit：{batch.limit}")
    lines.append("")

    lines.append("## 二、整体对比结果")
    lines.append("")
    lines.append("| 指标 | 数值 |")
    lines.append("| --- | --- |")
    lines.append(f"| 总样例数 | {batch.total_cases} |")
    lines.append(f"| 实际运行数 | {batch.ran} |")
    lines.append(f"| 失败数 | {batch.failed} |")
    lines.append(
        f"| RAG 平均耗时 | {format_report_number(batch.average_rag_latency_ms, ' ms')} |"
    )
    lines.append(
        f"| Agent 平均耗时 | {format_report_number(batch.average_agent_latency_ms, ' ms')} |"
    )
    lines.append(
        f"| RAG 平均 sources 数 | {format_report_number(batch.average_rag_sources_count)} |"
    )
    lines.append(
        f"| Agent 平均 sources 数 | {format_report_number(batch.average_agent_sources_count)} |"
    )
    lines.append(
        f"| Agent 平均工具调用数 | {format_report_number(batch.average_agent_tool_call_count)} |"
    )
    lines.append(f"| Agent 工具失败总数 | {batch.total_agent_failed_tool_count} |")
    lines.append("")

    lines.append("## 三、Agent 工具调用分析")
    lines.append("")
    lines.append("| 工具 | 调用次数 |")
    lines.append("| --- | --- |")
    lines.append(f"| summarize_document | {batch.summarize_document_count} |")
    lines.append(f"| read_document_chunks | {batch.read_document_chunks_count} |")
    lines.append(f"| search_rag_history | {batch.search_rag_history_count} |")
    lines.append("")

    lines.append("## 四、失败与异常样例分析")
    lines.append("")
    lines.append(f"- 失败/异常样例数：{len(failed_items)}")
    lines.append(f"- Agent 明显慢于 RAG 的样例数：{len(slow_agent_items)}")
    lines.append("")

    if not failed_items:
        lines.append("当前批次未发现明显失败或异常样例。")
        lines.append("")
    else:
        for index, item in enumerate(failed_items, start=1):
            reasons = build_item_warning_reasons(item)

            lines.append(f"### 失败样例 {index}")
            lines.append("")
            lines.append(f"- 问题：{item.question}")
            lines.append(f"- 异常原因：{'；'.join(reasons) if reasons else '未标明'}")
            lines.append(f"- RAG 耗时：{item.rag_latency_ms if item.rag_latency_ms is not None else '-'} ms")
            lines.append(f"- Agent 耗时：{item.agent_latency_ms if item.agent_latency_ms is not None else '-'} ms")
            lines.append(f"- RAG sources 数：{item.rag_sources_count}")
            lines.append(f"- Agent sources 数：{item.agent_sources_count}")
            lines.append(f"- Agent 工具链：{build_agent_tool_chain_text(item)}")

            if item.rag_error_message:
                lines.append(f"- RAG 错误：{item.rag_error_message}")

            if item.agent_error_message:
                lines.append(f"- Agent 错误：{item.agent_error_message}")

            failed_tool_calls = [
                tool_call
                for tool_call in item.agent_tool_calls
                if not tool_call.success
            ]

            if failed_tool_calls:
                lines.append("")
                lines.append("失败工具：")
                for tool_call in failed_tool_calls:
                    lines.append(
                        f"- {tool_call.tool_name}：{truncate_report_text(tool_call.observation, 300)}"
                    )

            lines.append("")

    lines.append("## 五、典型样例对比")
    lines.append("")
    lines.append("以下展示前 5 条样例的 RAG 与 Agent 回答对比。")
    lines.append("")

    for index, item in enumerate(items[:5], start=1):
        lines.append(f"### 样例 {index}")
        lines.append("")
        lines.append(f"**问题：** {item.question}")
        lines.append("")
        lines.append(f"**Agent 工具链：** {build_agent_tool_chain_text(item)}")
        lines.append("")
        lines.append("**RAG 回答：**")
        lines.append("")
        lines.append(truncate_report_text(item.rag_answer, 1000))
        lines.append("")
        lines.append("**Agent 回答：**")
        lines.append("")
        lines.append(truncate_report_text(item.agent_answer, 1000))
        lines.append("")
        lines.append("---")
        lines.append("")

    lines.append("## 六、实验结论")
    lines.append("")
    lines.append(
        "- 普通 RAG 主要体现一次检索与回答能力，适合直接知识问答场景。"
    )
    lines.append(
        "- Agent 通过工具调用链，可以执行文档总结、上下文扩展、历史检索等多步操作。"
    )
    lines.append(
        "- Agent 通常会带来更高的耗时成本，因此需要结合 max_steps、top_k 和工具失败率综合判断是否值得使用。"
    )
    lines.append(
        "- 如果 Agent 工具失败数较高，应优先检查工具参数解析、文档 chunk 是否存在、Planner 是否重复或误调用工具。"
    )
    lines.append("")
    lines.append("## 七、注意事项")
    lines.append("")
    lines.append(
        "本报告主要基于工程指标自动生成，包括耗时、sources 数量、工具调用数量和失败情况。"
    )
    lines.append(
        "当前报告尚未引入 LLM Judge 或人工评分，因此不能完全代表最终回答质量。"
    )

    return "\n".join(lines)

def markdown_cell(value: Any) -> str:
    if value is None:
        return "-"

    return str(value).replace("|", "\\|").replace("\n", " ").strip() or "-"


def report_percent(value: float | None) -> str:
    if value is None:
        return "-"

    return f"{value * 100:.1f}%"


def report_latency(value: float | int | None) -> str:
    if value is None:
        return "-"

    if value >= 1000:
        return f"{value / 1000:.2f}s"

    return f"{value:.0f}ms"


def report_number(value: float | int | None) -> str:
    if value is None:
        return "-"

    if isinstance(value, float):
        return f"{value:.2f}"

    return str(value)


def get_batch_failure_rate(batch: RagAgentCompareBatch) -> float | None:
    if not batch.ran:
        return None

    return batch.failed / batch.ran


def build_compare_batch_report_section(
    *,
    batch: RagAgentCompareBatch,
    title: str,
) -> str:
    failure_rate = get_batch_failure_rate(batch=batch)

    lines = [
        f"### {title}",
        "",
        f"- 批次名称：{batch.name}",
        f"- 批次 ID：{batch.id}",
        f"- 创建时间：{batch.created_at}",
        f"- 样例数量：{batch.total_cases}",
        f"- 成功运行：{batch.ran}",
        f"- 失败数量：{batch.failed}",
        f"- 失败率：{report_percent(failure_rate)}",
        (
            "- 参数："
            f"top_k={batch.top_k}，"
            f"max_steps={batch.max_steps}，"
            f"semantic_weight={batch.semantic_weight}，"
            f"keyword_weight={batch.keyword_weight}"
        ),
        f"- RAG 平均耗时：{report_latency(batch.average_rag_latency_ms)}",
        f"- Agent 平均耗时：{report_latency(batch.average_agent_latency_ms)}",
        f"- RAG 平均 sources 数：{report_number(batch.average_rag_sources_count)}",
        f"- Agent 平均 sources 数：{report_number(batch.average_agent_sources_count)}",
        f"- Agent 平均工具调用数：{report_number(batch.average_agent_tool_call_count)}",
        f"- Agent 工具失败总数：{batch.total_agent_failed_tool_count}",
        "",
    ]

    return "\n".join(lines)


def build_compare_batch_metric_rows(
    *,
    baseline_batch: RagAgentCompareBatch,
    optimized_batch: RagAgentCompareBatch,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    baseline_failure_rate = get_batch_failure_rate(batch=baseline_batch)
    optimized_failure_rate = get_batch_failure_rate(batch=optimized_batch)

    if baseline_failure_rate is not None and optimized_failure_rate is not None:
        change = optimized_failure_rate - baseline_failure_rate

        rows.append(
            {
                "metric": "失败率",
                "before": report_percent(baseline_failure_rate),
                "after": report_percent(optimized_failure_rate),
                "change": f"{change * 100:+.1f} 个百分点",
                "judgement": "改善" if change < 0 else "变差" if change > 0 else "基本不变",
            }
        )

    if (
        baseline_batch.average_agent_latency_ms is not None
        and optimized_batch.average_agent_latency_ms is not None
    ):
        change = (
            optimized_batch.average_agent_latency_ms
            - baseline_batch.average_agent_latency_ms
        )

        rows.append(
            {
                "metric": "Agent 平均耗时",
                "before": report_latency(baseline_batch.average_agent_latency_ms),
                "after": report_latency(optimized_batch.average_agent_latency_ms),
                "change": f"{change:+.0f}ms",
                "judgement": "改善" if change < 0 else "变差" if change > 0 else "基本不变",
            }
        )

    if (
        baseline_batch.average_agent_tool_call_count is not None
        and optimized_batch.average_agent_tool_call_count is not None
    ):
        change = (
            optimized_batch.average_agent_tool_call_count
            - baseline_batch.average_agent_tool_call_count
        )

        rows.append(
            {
                "metric": "Agent 平均工具调用数",
                "before": report_number(baseline_batch.average_agent_tool_call_count),
                "after": report_number(optimized_batch.average_agent_tool_call_count),
                "change": f"{change:+.2f}",
                "judgement": "观察",
            }
        )

    failed_tool_change = (
        optimized_batch.total_agent_failed_tool_count
        - baseline_batch.total_agent_failed_tool_count
    )

    rows.append(
        {
            "metric": "Agent 工具失败总数",
            "before": str(baseline_batch.total_agent_failed_tool_count),
            "after": str(optimized_batch.total_agent_failed_tool_count),
            "change": f"{failed_tool_change:+d}",
            "judgement": "改善"
            if failed_tool_change < 0
            else "变差"
            if failed_tool_change > 0
            else "基本不变",
        }
    )

    if (
        baseline_batch.average_rag_sources_count is not None
        and baseline_batch.average_agent_sources_count is not None
        and optimized_batch.average_rag_sources_count is not None
        and optimized_batch.average_agent_sources_count is not None
    ):
        baseline_source_gain = (
            baseline_batch.average_agent_sources_count
            - baseline_batch.average_rag_sources_count
        )

        optimized_source_gain = (
            optimized_batch.average_agent_sources_count
            - optimized_batch.average_rag_sources_count
        )

        change = optimized_source_gain - baseline_source_gain

        rows.append(
            {
                "metric": "Agent 相对 RAG 的 sources 增益",
                "before": report_number(baseline_source_gain),
                "after": report_number(optimized_source_gain),
                "change": f"{change:+.2f}",
                "judgement": "改善" if change > 0 else "变差" if change < 0 else "基本不变",
            }
        )

    return rows


def build_code_eval_type_stats_for_compare_batch(
    *,
    session: Session,
    batch_id: uuid.UUID,
) -> list[dict[str, Any]]:
    items = session.exec(
        select(RagAgentCompareItem)
        .where(RagAgentCompareItem.batch_id == batch_id)
        .order_by(col(RagAgentCompareItem.created_at).asc())
    ).all()

    eval_case_ids = [
        item.eval_case_id
        for item in items
        if item.eval_case_id is not None
    ]

    eval_cases_by_id: dict[uuid.UUID, RagEvalCase] = {}

    if eval_case_ids:
        eval_cases = session.exec(
            select(RagEvalCase).where(col(RagEvalCase.id).in_(eval_case_ids))
        ).all()

        eval_cases_by_id = {
            eval_case.id: eval_case
            for eval_case in eval_cases
        }

    stat_map: dict[str, dict[str, Any]] = {}

    for item in items:
        eval_case = None

        if item.eval_case_id is not None:
            eval_case = eval_cases_by_id.get(item.eval_case_id)

        case_type = "unknown"

        if eval_case is not None:
            case_type = parse_code_eval_case_type(eval_case.note)

        if case_type not in stat_map:
            stat_map[case_type] = {
                "case_type": case_type,
                "case_type_label": get_code_eval_case_type_label(case_type),
                "compared_cases": 0,
                "rag_source_hit_count": 0,
                "agent_source_hit_count": 0,
                "source_check_count": 0,
                "rag_keyword_hit_sum": 0.0,
                "agent_keyword_hit_sum": 0.0,
                "keyword_check_count": 0,
                "rag_latency_sum": 0,
                "agent_latency_sum": 0,
                "latency_count": 0,
                "agent_tool_call_sum": 0,
            }

        stat = stat_map[case_type]
        stat["compared_cases"] += 1

        if item.expected_source_filename:
            stat["source_check_count"] += 1

            if calculate_source_file_hit(
                sources_value=item.rag_sources_json,
                expected_source_filename=item.expected_source_filename,
            ):
                stat["rag_source_hit_count"] += 1

            if calculate_source_file_hit(
                sources_value=item.agent_sources_json,
                expected_source_filename=item.expected_source_filename,
            ):
                stat["agent_source_hit_count"] += 1

        expected_keywords = safe_json_loads(item.expected_keywords_json)

        if isinstance(expected_keywords, list) and expected_keywords:
            stat["keyword_check_count"] += 1
            stat["rag_keyword_hit_sum"] += calculate_keyword_hit_rate(
                answer=item.rag_answer,
                expected_keywords_json=item.expected_keywords_json,
            )
            stat["agent_keyword_hit_sum"] += calculate_keyword_hit_rate(
                answer=item.agent_answer,
                expected_keywords_json=item.expected_keywords_json,
            )

        if item.rag_latency_ms is not None and item.agent_latency_ms is not None:
            stat["latency_count"] += 1
            stat["rag_latency_sum"] += item.rag_latency_ms
            stat["agent_latency_sum"] += item.agent_latency_ms

        stat["agent_tool_call_sum"] += item.agent_tool_call_count

    result: list[dict[str, Any]] = []

    for stat in stat_map.values():
        source_count = stat["source_check_count"]
        keyword_count = stat["keyword_check_count"]
        latency_count = stat["latency_count"]
        compared_cases = stat["compared_cases"]

        result.append(
            {
                "case_type": stat["case_type"],
                "case_type_label": stat["case_type_label"],
                "compared_cases": compared_cases,
                "rag_source_hit_rate": (
                    stat["rag_source_hit_count"] / source_count
                    if source_count
                    else None
                ),
                "agent_source_hit_rate": (
                    stat["agent_source_hit_count"] / source_count
                    if source_count
                    else None
                ),
                "rag_keyword_hit_rate": (
                    stat["rag_keyword_hit_sum"] / keyword_count
                    if keyword_count
                    else None
                ),
                "agent_keyword_hit_rate": (
                    stat["agent_keyword_hit_sum"] / keyword_count
                    if keyword_count
                    else None
                ),
                "rag_avg_latency_ms": (
                    stat["rag_latency_sum"] / latency_count
                    if latency_count
                    else None
                ),
                "agent_avg_latency_ms": (
                    stat["agent_latency_sum"] / latency_count
                    if latency_count
                    else None
                ),
                "agent_avg_tool_calls": (
                    stat["agent_tool_call_sum"] / compared_cases
                    if compared_cases
                    else None
                ),
            }
        )

    return result


def build_compare_failure_report_data(
    *,
    session: Session,
    batch_id: uuid.UUID,
    limit: int = 8,
) -> tuple[dict[str, int], list[RagAgentCompareItem]]:
    items = session.exec(
        select(RagAgentCompareItem)
        .where(RagAgentCompareItem.batch_id == batch_id)
        .order_by(col(RagAgentCompareItem.created_at).desc())
    ).all()

    reason_count_map: dict[str, int] = {}
    failure_items: list[RagAgentCompareItem] = []

    for item in items:
        reasons = build_compare_failure_reasons(item=item)

        if not reasons:
            continue

        for reason in reasons:
            reason_count_map[reason] = reason_count_map.get(reason, 0) + 1

        if len(failure_items) < limit:
            failure_items.append(item)

    return reason_count_map, failure_items


def build_backend_code_agent_advice(
    *,
    reason_count_map: dict[str, int],
    type_stats: list[dict[str, Any]],
) -> list[str]:
    advice: list[str] = []

    if reason_count_map.get("agent_no_tool_calls", 0) > 0:
        advice.append(
            "Agent 存在无工具调用情况，建议强化 Planner prompt，让代码类问题优先调用 search_code、find_code_references 或 read_code_file。"
        )

    if reason_count_map.get("agent_tool_failed", 0) > 0:
        advice.append(
            "Agent 存在工具调用失败，建议优先检查 agent_tool_registry 中的 handler 参数解析、异常处理和返回结构。"
        )

    if reason_count_map.get("agent_source_miss", 0) > 0:
        advice.append(
            "Agent 来源文件未命中，建议提高 keyword_weight，例如尝试 semantic_weight=0.65、keyword_weight=0.35。"
        )

    if reason_count_map.get("agent_keyword_miss", 0) > 0:
        advice.append(
            "Agent 回答关键词未命中，建议在最终回答 prompt 中要求明确输出文件名、函数名、类名、line_range 和核心逻辑。"
        )

    if reason_count_map.get("agent_slow", 0) > 0:
        advice.append(
            "Agent 明显慢于 RAG，建议限制重复工具调用，尤其是 read_code_file 和 search_code 的重复调用。"
        )

    for stat in type_stats:
        case_type = stat["case_type"]

        if (
            case_type == "symbol_reference"
            and stat["agent_source_hit_rate"] is not None
            and stat["agent_source_hit_rate"] < 0.7
        ):
            advice.append(
                "symbol_reference 类型任务的 Agent 来源命中率偏低，建议让“在哪里调用、哪里使用、引用位置、调用关系”类问题优先触发 find_code_references。"
            )

        if (
            case_type == "code_logic"
            and stat["agent_keyword_hit_rate"] is not None
            and stat["agent_keyword_hit_rate"] < 0.7
        ):
            advice.append(
                "code_logic 类型任务的 Agent 关键词命中率偏低，建议在 search_code 后继续 read_code_file，结合完整文件上下文回答。"
            )

    if not advice:
        advice.append(
            "当前未发现突出的异常模式。建议扩大评测集规模，并加入 bug_location、跨文件调用链分析等更难任务。"
        )

    return advice


def build_code_agent_experiment_report_markdown(
    *,
    knowledge_base: KnowledgeBase,
    baseline_batch: RagAgentCompareBatch,
    optimized_batch: RagAgentCompareBatch,
    metric_rows: list[dict[str, str]],
    type_stats: list[dict[str, Any]],
    reason_count_map: dict[str, int],
    failure_items: list[RagAgentCompareItem],
    advice: list[str],
) -> str:
    lines: list[str] = []

    better_count = sum(1 for row in metric_rows if row["judgement"] == "改善")
    worse_count = sum(1 for row in metric_rows if row["judgement"] == "变差")

    if better_count > worse_count:
        auto_conclusion = "本轮优化整体呈现正向效果，改善项多于变差项。"
    elif worse_count > better_count:
        auto_conclusion = "本轮优化存在一定副作用，建议结合失败案例继续定位问题。"
    else:
        auto_conclusion = "本轮优化效果暂不明显，建议扩大评测样例并继续优化。"

    lines.extend(
        [
            "# Code RAG / Code Agent 实验报告",
            "",
            f"生成时间：{datetime.now().isoformat(timespec='seconds')}",
            "",
            "## 1. 实验对象",
            "",
            f"- 知识库名称：{knowledge_base.name}",
            f"- 知识库 ID：{knowledge_base.id}",
            f"- 知识库描述：{knowledge_base.description or '-'}",
            "",
            "## 2. 实验批次信息",
            "",
            build_compare_batch_report_section(
                batch=baseline_batch,
                title="优化前批次",
            ),
            build_compare_batch_report_section(
                batch=optimized_batch,
                title="优化后批次",
            ),
            "## 3. 优化前后总体对比",
            "",
            "| 指标 | 优化前 | 优化后 | 变化 | 判断 |",
            "|---|---:|---:|---:|---|",
        ]
    )

    for row in metric_rows:
        lines.append(
            "| "
            f"{markdown_cell(row['metric'])} | "
            f"{markdown_cell(row['before'])} | "
            f"{markdown_cell(row['after'])} | "
            f"{markdown_cell(row['change'])} | "
            f"{markdown_cell(row['judgement'])} |"
        )

    lines.extend(["", f"自动结论：{auto_conclusion}", ""])

    lines.extend(
        [
            "## 4. 分任务类型统计",
            "",
            "| 任务类型 | 已对比 | RAG 文件命中 | Agent 文件命中 | RAG 关键词命中 | Agent 关键词命中 | RAG 耗时 | Agent 耗时 | Agent 平均工具数 |",
            "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
        ]
    )

    if not type_stats:
        lines.append("| 暂无数据 | - | - | - | - | - | - | - | - |")
    else:
        for stat in type_stats:
            lines.append(
                "| "
                f"{markdown_cell(stat['case_type_label'])} | "
                f"{stat['compared_cases']} | "
                f"{report_percent(stat['rag_source_hit_rate'])} | "
                f"{report_percent(stat['agent_source_hit_rate'])} | "
                f"{report_percent(stat['rag_keyword_hit_rate'])} | "
                f"{report_percent(stat['agent_keyword_hit_rate'])} | "
                f"{report_latency(stat['rag_avg_latency_ms'])} | "
                f"{report_latency(stat['agent_avg_latency_ms'])} | "
                f"{report_number(stat['agent_avg_tool_calls'])} |"
            )

    lines.extend(
        [
            "",
            "## 5. 失败原因统计",
            "",
            "| 失败原因 | 次数 |",
            "|---|---:|",
        ]
    )

    if not reason_count_map:
        lines.append("| 暂无明显失败原因 | - |")
    else:
        for reason, count in sorted(
            reason_count_map.items(),
            key=lambda item: item[1],
            reverse=True,
        ):
            lines.append(
                f"| {markdown_cell(get_compare_failure_reason_label(reason))} | {count} |"
            )

    lines.extend(["", "## 6. 代表性失败案例", ""])

    if not failure_items:
        lines.append("当前批次暂无失败案例。")
    else:
        for index, item in enumerate(failure_items, start=1):
            reasons = build_compare_failure_reasons(item=item)
            reason_labels = [
                get_compare_failure_reason_label(reason)
                for reason in reasons
            ]

            lines.extend(
                [
                    f"### 6.{index} {item.question}",
                    "",
                    f"- eval_case_id：{item.eval_case_id or '-'}",
                    f"- 失败原因：{'、'.join(reason_labels) if reason_labels else '-'}",
                    f"- RAG 耗时：{report_latency(item.rag_latency_ms)}",
                    f"- Agent 耗时：{report_latency(item.agent_latency_ms)}",
                    f"- Agent 工具调用数：{item.agent_tool_call_count}",
                    f"- Agent 工具失败数：{item.agent_failed_tool_count}",
                    "",
                ]
            )

    lines.extend(["", "## 7. 自动优化建议", ""])

    for index, item in enumerate(advice, start=1):
        lines.append(f"{index}. {item}")

    lines.extend(
        [
            "",
            "## 8. 下一轮实验建议",
            "",
            "- 如果 Agent 来源未命中较多，建议尝试 semantic_weight=0.65、keyword_weight=0.35。",
            "- 如果 symbol_reference 表现较差，建议强化 find_code_references 的触发规则。",
            "- 如果 code_logic 表现较差，建议 search_code 后继续 read_code_file。",
            "- 如果 Agent 耗时明显升高，建议限制重复工具调用。",
            "",
        ]
    )

    return "\n".join(lines)


def set_docx_run_font(run, font_name: str = "宋体", font_size: int = 10) -> None:
    run.font.name = font_name
    run.font.size = Pt(font_size)

    run_element = run._element
    run_properties = run_element.get_or_add_rPr()
    run_properties.rFonts.set(qn("w:eastAsia"), font_name)


def add_docx_paragraph(
    *,
    document: DocxDocument,
    text: str,
    style: str | None = None,
    font_size: int = 10,
) -> None:
    paragraph = document.add_paragraph(style=style)

    for index, line in enumerate(text.split("\n")):
        if index > 0:
            paragraph.add_run().add_break()

        run = paragraph.add_run(line)
        set_docx_run_font(run, font_size=font_size)


def add_docx_metric_table(
    *,
    document: DocxDocument,
    rows: list[tuple[str, str]],
) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    header_cells[0].text = "指标"
    header_cells[1].text = "数值"

    for metric, value in rows:
        cells = table.add_row().cells
        cells[0].text = metric
        cells[1].text = value

    document.add_paragraph()


def build_rag_agent_compare_docx_report(
    *,
    knowledge_base_name: str,
    batch: RagAgentCompareBatchPublic,
    items: list[RagAgentCompareItemPublic],
) -> bytes:
    failed_items = [
        item for item in items
        if item.is_failed
        or item.rag_error_message
        or item.agent_error_message
        or item.agent_failed_tool_count > 0
        or item.rag_sources_count == 0
        or item.agent_sources_count == 0
    ]

    slow_agent_items = [
        item for item in items
        if item.rag_latency_ms is not None
        and item.agent_latency_ms is not None
        and (
            item.agent_latency_ms > item.rag_latency_ms * 2
            or item.agent_latency_ms - item.rag_latency_ms > 3000
        )
    ]

    document = DocxDocument()

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "宋体"
    normal_style.font.size = Pt(10)
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title = document.add_heading("RAG vs Agent 实验报告", level=0)
    for run in title.runs:
        set_docx_run_font(run, font_name="宋体", font_size=18)

    document.add_heading("一、实验基本信息", level=1)

    basic_rows = [
        ("知识库", knowledge_base_name),
        ("批次名称", batch.name),
        ("批次 ID", str(batch.id)),
        ("创建时间", batch.created_at.isoformat() if batch.created_at else "-"),
        ("top_k", str(batch.top_k)),
        ("max_steps", str(batch.max_steps)),
        ("semantic_weight", str(batch.semantic_weight)),
        ("keyword_weight", str(batch.keyword_weight)),
        ("limit", str(batch.limit)),
    ]
    add_docx_metric_table(document=document, rows=basic_rows)

    document.add_heading("二、整体对比结果", level=1)

    summary_rows = [
        ("总样例数", str(batch.total_cases)),
        ("实际运行数", str(batch.ran)),
        ("失败数", str(batch.failed)),
        ("RAG 平均耗时", format_report_number(batch.average_rag_latency_ms, " ms")),
        ("Agent 平均耗时", format_report_number(batch.average_agent_latency_ms, " ms")),
        ("RAG 平均 sources 数", format_report_number(batch.average_rag_sources_count)),
        ("Agent 平均 sources 数", format_report_number(batch.average_agent_sources_count)),
        ("Agent 平均工具调用数", format_report_number(batch.average_agent_tool_call_count)),
        ("Agent 工具失败总数", str(batch.total_agent_failed_tool_count)),
    ]
    add_docx_metric_table(document=document, rows=summary_rows)

    document.add_heading("三、Agent 工具调用分析", level=1)

    tool_rows = [
        ("summarize_document", str(batch.summarize_document_count)),
        ("read_document_chunks", str(batch.read_document_chunks_count)),
        ("search_rag_history", str(batch.search_rag_history_count)),
    ]
    add_docx_metric_table(document=document, rows=tool_rows)

    document.add_heading("四、失败与异常样例分析", level=1)

    add_docx_paragraph(
        document=document,
        text=(
            f"失败/异常样例数：{len(failed_items)}\n"
            f"Agent 明显慢于 RAG 的样例数：{len(slow_agent_items)}"
        ),
    )

    if not failed_items:
        add_docx_paragraph(
            document=document,
            text="当前批次未发现明显失败或异常样例。",
        )
    else:
        for index, item in enumerate(failed_items, start=1):
            document.add_heading(f"失败样例 {index}", level=2)

            reasons = build_item_warning_reasons(item)

            add_docx_paragraph(
                document=document,
                text=(
                    f"问题：{item.question}\n"
                    f"异常原因：{'；'.join(reasons) if reasons else '未标明'}\n"
                    f"RAG 耗时：{item.rag_latency_ms if item.rag_latency_ms is not None else '-'} ms\n"
                    f"Agent 耗时：{item.agent_latency_ms if item.agent_latency_ms is not None else '-'} ms\n"
                    f"RAG sources 数：{item.rag_sources_count}\n"
                    f"Agent sources 数：{item.agent_sources_count}\n"
                    f"Agent 工具链：{build_agent_tool_chain_text(item)}"
                ),
            )

            if item.rag_error_message:
                add_docx_paragraph(
                    document=document,
                    text=f"RAG 错误：{item.rag_error_message}",
                )

            if item.agent_error_message:
                add_docx_paragraph(
                    document=document,
                    text=f"Agent 错误：{item.agent_error_message}",
                )

    document.add_heading("五、典型样例对比", level=1)

    for index, item in enumerate(items[:5], start=1):
        document.add_heading(f"样例 {index}", level=2)

        add_docx_paragraph(
            document=document,
            text=(
                f"问题：{item.question}\n"
                f"Agent 工具链：{build_agent_tool_chain_text(item)}"
            ),
        )

        document.add_heading("RAG 回答", level=3)
        add_docx_paragraph(
            document=document,
            text=truncate_report_text(item.rag_answer, 1000),
        )

        document.add_heading("Agent 回答", level=3)
        add_docx_paragraph(
            document=document,
            text=truncate_report_text(item.agent_answer, 1000),
        )

    document.add_heading("六、实验结论", level=1)

    add_docx_paragraph(
        document=document,
        text=(
            "普通 RAG 主要体现一次检索与回答能力，适合直接知识问答场景。\n"
            "Agent 通过工具调用链，可以执行文档总结、上下文扩展、历史检索等多步操作。\n"
            "Agent 通常会带来更高的耗时成本，因此需要结合 max_steps、top_k 和工具失败率综合判断是否值得使用。\n"
            "如果 Agent 工具失败数较高，应优先检查工具参数解析、文档 chunk 是否存在、Planner 是否重复或误调用工具。"
        ),
    )

    document.add_heading("七、注意事项", level=1)

    add_docx_paragraph(
        document=document,
        text=(
            "本报告主要基于工程指标自动生成，包括耗时、sources 数量、工具调用数量和失败情况。\n"
            "当前报告尚未引入 LLM Judge 或人工评分，因此不能完全代表最终回答质量。"
        ),
    )

    output = BytesIO()
    document.save(output)

    return output.getvalue()

def register_pdf_chinese_font() -> str:
    font_name = "STSong-Light"

    try:
        pdfmetrics.getFont(font_name)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))

    return font_name


def pdf_paragraph(
    text: str | None,
    style: ParagraphStyle,
) -> Paragraph:
    safe_text = escape(text or "无")
    safe_text = safe_text.replace("\n", "<br/>")

    return Paragraph(safe_text, style)


def build_pdf_table(
    *,
    rows: list[list[str]],
    style: ParagraphStyle,
    col_widths: list[float] | None = None,
) -> Table:
    table_data = [
        [pdf_paragraph(cell, style) for cell in row]
        for row in rows
    ]

    table = Table(
        table_data,
        colWidths=col_widths,
        hAlign="LEFT",
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )

    return table


def build_rag_agent_compare_pdf_report(
    *,
    knowledge_base_name: str,
    batch: RagAgentCompareBatchPublic,
    items: list[RagAgentCompareItemPublic],
) -> bytes:
    font_name = register_pdf_chinese_font()

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=24,
        alignment=TA_LEFT,
        spaceAfter=14,
    )

    heading1_style = ParagraphStyle(
        "ChineseHeading1",
        parent=styles["Heading1"],
        fontName=font_name,
        fontSize=14,
        leading=20,
        spaceBefore=12,
        spaceAfter=8,
    )

    heading2_style = ParagraphStyle(
        "ChineseHeading2",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=12,
        leading=18,
        spaceBefore=10,
        spaceAfter=6,
    )

    normal_style = ParagraphStyle(
        "ChineseNormal",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=14,
        spaceAfter=6,
    )

    table_style = ParagraphStyle(
        "ChineseTable",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8,
        leading=12,
    )

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title="RAG vs Agent 实验报告",
    )

    story = []

    story.append(Paragraph("RAG vs Agent 实验报告", title_style))

    story.append(Paragraph("一、实验基本信息", heading1_style))

    basic_rows = [
        ["指标", "数值"],
        ["知识库", knowledge_base_name],
        ["批次名称", batch.name],
        ["批次 ID", str(batch.id)],
        ["创建时间", batch.created_at.isoformat() if batch.created_at else "-"],
        ["top_k", str(batch.top_k)],
        ["max_steps", str(batch.max_steps)],
        ["semantic_weight", str(batch.semantic_weight)],
        ["keyword_weight", str(batch.keyword_weight)],
        ["limit", str(batch.limit)],
    ]

    story.append(
        build_pdf_table(
            rows=basic_rows,
            style=table_style,
            col_widths=[4 * cm, 12 * cm],
        )
    )
    story.append(Spacer(1, 10))

    story.append(Paragraph("二、整体对比结果", heading1_style))

    summary_rows = [
        ["指标", "数值"],
        ["总样例数", str(batch.total_cases)],
        ["实际运行数", str(batch.ran)],
        ["失败数", str(batch.failed)],
        ["RAG 平均耗时", format_report_number(batch.average_rag_latency_ms, " ms")],
        ["Agent 平均耗时", format_report_number(batch.average_agent_latency_ms, " ms")],
        ["RAG 平均 sources 数", format_report_number(batch.average_rag_sources_count)],
        ["Agent 平均 sources 数", format_report_number(batch.average_agent_sources_count)],
        ["Agent 平均工具调用数", format_report_number(batch.average_agent_tool_call_count)],
        ["Agent 工具失败总数", str(batch.total_agent_failed_tool_count)],
    ]

    story.append(
        build_pdf_table(
            rows=summary_rows,
            style=table_style,
            col_widths=[6 * cm, 10 * cm],
        )
    )
    story.append(Spacer(1, 10))

    story.append(Paragraph("三、Agent 工具调用分析", heading1_style))

    tool_rows = [
        ["工具", "调用次数"],
        ["summarize_document", str(batch.summarize_document_count)],
        ["read_document_chunks", str(batch.read_document_chunks_count)],
        ["search_rag_history", str(batch.search_rag_history_count)],
    ]

    story.append(
        build_pdf_table(
            rows=tool_rows,
            style=table_style,
            col_widths=[8 * cm, 8 * cm],
        )
    )
    story.append(Spacer(1, 10))

    failed_items = [
        item for item in items
        if item.is_failed
        or item.rag_error_message
        or item.agent_error_message
        or item.agent_failed_tool_count > 0
        or item.rag_sources_count == 0
        or item.agent_sources_count == 0
    ]

    slow_agent_items = [
        item for item in items
        if item.rag_latency_ms is not None
        and item.agent_latency_ms is not None
        and (
            item.agent_latency_ms > item.rag_latency_ms * 2
            or item.agent_latency_ms - item.rag_latency_ms > 3000
        )
    ]

    story.append(Paragraph("四、失败与异常样例分析", heading1_style))
    story.append(
        pdf_paragraph(
            (
                f"失败/异常样例数：{len(failed_items)}\n"
                f"Agent 明显慢于 RAG 的样例数：{len(slow_agent_items)}"
            ),
            normal_style,
        )
    )

    if not failed_items:
        story.append(pdf_paragraph("当前批次未发现明显失败或异常样例。", normal_style))
    else:
        for index, item in enumerate(failed_items[:10], start=1):
            story.append(Paragraph(f"失败样例 {index}", heading2_style))

            reasons = build_item_warning_reasons(item)

            story.append(
                pdf_paragraph(
                    (
                        f"问题：{item.question}\n"
                        f"异常原因：{'；'.join(reasons) if reasons else '未标明'}\n"
                        f"RAG 耗时：{item.rag_latency_ms if item.rag_latency_ms is not None else '-'} ms\n"
                        f"Agent 耗时：{item.agent_latency_ms if item.agent_latency_ms is not None else '-'} ms\n"
                        f"RAG sources 数：{item.rag_sources_count}\n"
                        f"Agent sources 数：{item.agent_sources_count}\n"
                        f"Agent 工具链：{build_agent_tool_chain_text(item)}"
                    ),
                    normal_style,
                )
            )

    story.append(PageBreak())

    story.append(Paragraph("五、典型样例对比", heading1_style))

    for index, item in enumerate(items[:5], start=1):
        story.append(Paragraph(f"样例 {index}", heading2_style))
        story.append(pdf_paragraph(f"问题：{item.question}", normal_style))
        story.append(pdf_paragraph(f"Agent 工具链：{build_agent_tool_chain_text(item)}", normal_style))

        story.append(Paragraph("RAG 回答", heading2_style))
        story.append(pdf_paragraph(truncate_report_text(item.rag_answer, 800), normal_style))

        story.append(Paragraph("Agent 回答", heading2_style))
        story.append(pdf_paragraph(truncate_report_text(item.agent_answer, 800), normal_style))

        story.append(Spacer(1, 8))

    story.append(Paragraph("六、实验结论", heading1_style))
    story.append(
        pdf_paragraph(
            (
                "普通 RAG 主要体现一次检索与回答能力，适合直接知识问答场景。\n"
                "Agent 通过工具调用链，可以执行文档总结、上下文扩展、历史检索等多步操作。\n"
                "Agent 通常会带来更高的耗时成本，因此需要结合 max_steps、top_k 和工具失败率综合判断是否值得使用。\n"
                "如果 Agent 工具失败数较高，应优先检查工具参数解析、文档 chunk 是否存在、Planner 是否重复或误调用工具。"
            ),
            normal_style,
        )
    )

    story.append(Paragraph("七、注意事项", heading1_style))
    story.append(
        pdf_paragraph(
            (
                "本报告主要基于工程指标自动生成，包括耗时、sources 数量、工具调用数量和失败情况。\n"
                "当前报告尚未引入 LLM Judge 或人工评分，因此不能完全代表最终回答质量。"
            ),
            normal_style,
        )
    )

    doc.build(story)

    return buffer.getvalue()


def parse_json_list(value: str | None) -> list[str]:
    if not value:
        return []

    try:
        data = json.loads(value)
    except json.JSONDecodeError:
        return []

    if not isinstance(data, list):
        return []

    return [str(item) for item in data]

def parse_code_eval_case_type(note: str | None) -> str:
    if not note:
        return "unknown"

    if note.startswith("code:"):
        return note.replace("code:", "", 1)

    if "代码文件列表" in note or "代码文件作用" in note:
        return "file_overview"

    if "定义与作用" in note:
        return "symbol_definition"

    if "引用定位" in note:
        return "symbol_reference"

    return "unknown"


def get_code_eval_case_type_label(case_type: str) -> str:
    label_map = {
        "file_overview": "文件概览",
        "symbol_definition": "符号定义",
        "symbol_reference": "引用定位",
        "code_logic": "代码逻辑",
        "bug_location": "Bug 定位",
        "unknown": "未分类",
    }

    return label_map.get(case_type, case_type)

def safe_json_loads(value):
    if value is None:
        return None

    if isinstance(value, (list, dict)):
        return value

    if not isinstance(value, str):
        return None

    value = value.strip()

    if not value:
        return None

    try:
        return json.loads(value)
    except json.JSONDecodeError:
        return None


def calculate_keyword_hit_rate(
    *,
    answer: str | None,
    expected_keywords_json: str | None,
) -> float:
    if not answer:
        return 0

    expected_keywords = safe_json_loads(expected_keywords_json)

    if not expected_keywords or not isinstance(expected_keywords, list):
        return 0

    valid_keywords = [
        str(keyword).strip()
        for keyword in expected_keywords
        if str(keyword).strip()
    ]

    if not valid_keywords:
        return 0

    answer_lower = answer.lower()

    hit_count = 0

    for keyword in valid_keywords:
        if keyword.lower() in answer_lower:
            hit_count += 1

    return hit_count / len(valid_keywords)


def calculate_source_file_hit(
    *,
    sources_value,
    expected_source_filename: str | None,
) -> bool:
    if not expected_source_filename:
        return False

    sources = safe_json_loads(sources_value)

    if not sources or not isinstance(sources, list):
        return False

    expected_name = expected_source_filename.lower()

    for source in sources:
        if not isinstance(source, dict):
            continue

        filename = (
            source.get("original_filename")
            or source.get("filename")
            or source.get("source_filename")
            or ""
        )

        if expected_name in str(filename).lower():
            return True

    return False


def get_compare_failure_reason_label(reason: str) -> str:
    label_map = {
        "rag_error": "RAG 运行出错",
        "agent_error": "Agent 运行出错",
        "rag_no_sources": "RAG 无 sources",
        "agent_no_sources": "Agent 无 sources",
        "agent_no_tool_calls": "Agent 无工具调用",
        "agent_tool_failed": "Agent 工具调用失败",
        "rag_source_miss": "RAG 来源未命中",
        "agent_source_miss": "Agent 来源未命中",
        "rag_keyword_miss": "RAG 关键词未命中",
        "agent_keyword_miss": "Agent 关键词未命中",
        "agent_slow": "Agent 明显慢于 RAG",
        "compare_marked_failed": "对比项被标记失败",
    }

    return label_map.get(reason, reason)


def build_compare_failure_reasons(
    *,
    item: RagAgentCompareItem,
) -> list[str]:
    reasons: list[str] = []

    expected_keywords_json = item.expected_keywords_json
    expected_source_filename = item.expected_source_filename

    rag_keyword_hit_rate = calculate_keyword_hit_rate(
        answer=item.rag_answer,
        expected_keywords_json=expected_keywords_json,
    )

    agent_keyword_hit_rate = calculate_keyword_hit_rate(
        answer=item.agent_answer,
        expected_keywords_json=expected_keywords_json,
    )

    rag_source_hit = calculate_source_file_hit(
        sources_value=item.rag_sources_json,
        expected_source_filename=expected_source_filename,
    )

    agent_source_hit = calculate_source_file_hit(
        sources_value=item.agent_sources_json,
        expected_source_filename=expected_source_filename,
    )

    if item.is_failed:
        reasons.append("compare_marked_failed")

    if item.rag_error_message:
        reasons.append("rag_error")

    if item.agent_error_message:
        reasons.append("agent_error")

    if item.rag_sources_count == 0:
        reasons.append("rag_no_sources")

    if item.agent_sources_count == 0:
        reasons.append("agent_no_sources")

    if item.agent_tool_call_count == 0:
        reasons.append("agent_no_tool_calls")

    if item.agent_failed_tool_count > 0:
        reasons.append("agent_tool_failed")

    if expected_source_filename:
        if not rag_source_hit:
            reasons.append("rag_source_miss")

        if not agent_source_hit:
            reasons.append("agent_source_miss")

    expected_keywords = safe_json_loads(expected_keywords_json)

    if isinstance(expected_keywords, list) and len(expected_keywords) > 0:
        if rag_keyword_hit_rate < 1:
            reasons.append("rag_keyword_miss")

        if agent_keyword_hit_rate < 1:
            reasons.append("agent_keyword_miss")

    if (
        item.rag_latency_ms is not None
        and item.agent_latency_ms is not None
        and (
            item.agent_latency_ms > item.rag_latency_ms * 2
            or item.agent_latency_ms - item.rag_latency_ms > 3000
        )
    ):
        reasons.append("agent_slow")

    return reasons


def build_compare_failure_case_public(
    *,
    item: RagAgentCompareItem,
    eval_case: RagEvalCase | None,
) -> RagAgentCompareFailureCasePublic:
    case_type = "unknown"

    if eval_case is not None:
        case_type = parse_code_eval_case_type(eval_case.note)

    failure_reasons = build_compare_failure_reasons(item=item)

    rag_source_hit = calculate_source_file_hit(
        sources_value=item.rag_sources_json,
        expected_source_filename=item.expected_source_filename,
    )

    agent_source_hit = calculate_source_file_hit(
        sources_value=item.agent_sources_json,
        expected_source_filename=item.expected_source_filename,
    )

    rag_keyword_hit_rate = calculate_keyword_hit_rate(
        answer=item.rag_answer,
        expected_keywords_json=item.expected_keywords_json,
    )

    agent_keyword_hit_rate = calculate_keyword_hit_rate(
        answer=item.agent_answer,
        expected_keywords_json=item.expected_keywords_json,
    )

    return RagAgentCompareFailureCasePublic(
        id=item.id,
        batch_id=item.batch_id,
        eval_case_id=item.eval_case_id,
        question=item.question,
        case_type=case_type,
        case_type_label=get_code_eval_case_type_label(case_type),
        failure_reasons=failure_reasons,
        failure_reason_labels=[
            get_compare_failure_reason_label(reason)
            for reason in failure_reasons
        ],
        rag_source_hit=rag_source_hit,
        agent_source_hit=agent_source_hit,
        rag_keyword_hit_rate=rag_keyword_hit_rate,
        agent_keyword_hit_rate=agent_keyword_hit_rate,
        rag_latency_ms=item.rag_latency_ms,
        agent_latency_ms=item.agent_latency_ms,
        rag_sources_count=item.rag_sources_count,
        agent_sources_count=item.agent_sources_count,
        agent_tool_call_count=item.agent_tool_call_count,
        agent_failed_tool_count=item.agent_failed_tool_count,
        agent_tool_names=parse_json_data(
            item.agent_tool_names_json,
            [],
        ),
        rag_error_message=item.rag_error_message,
        agent_error_message=item.agent_error_message,
        rag_answer=item.rag_answer,
        agent_answer=item.agent_answer,
        created_at=item.created_at,
    )


def count_agent_tool_calls(tool_calls_value) -> tuple[int, dict[str, int]]:
    tool_calls = safe_json_loads(tool_calls_value)

    if not tool_calls or not isinstance(tool_calls, list):
        return 0, {}

    usage: dict[str, int] = {}

    for tool_call in tool_calls:
        if not isinstance(tool_call, dict):
            continue

        tool_name = (
            tool_call.get("tool_name")
            or tool_call.get("name")
            or tool_call.get("action")
            or "unknown"
        )

        tool_name = str(tool_name)

        usage[tool_name] = usage.get(tool_name, 0) + 1

    return len(tool_calls), usage


def merge_tool_usage(
    *,
    target: dict[str, int],
    current: dict[str, int],
) -> None:
    for tool_name, count in current.items():
        target[tool_name] = target.get(tool_name, 0) + count


def rag_eval_case_to_public(eval_case: RagEvalCase) -> RagEvalCasePublic:
    return RagEvalCasePublic(
        id=eval_case.id,
        knowledge_base_id=eval_case.knowledge_base_id,
        owner_id=eval_case.owner_id,
        question=eval_case.question,
        expected_keywords=parse_json_list(eval_case.expected_keywords_json),
        expected_source_filename=eval_case.expected_source_filename,
        note=eval_case.note,
        created_at=eval_case.created_at,
    )


def rag_eval_run_to_public(eval_run: RagEvalRun) -> RagEvalRunPublic:
    try:
        source_items = json.loads(eval_run.sources_json or "[]")
    except json.JSONDecodeError:
        source_items = []

    sources = []

    for item in source_items:
        try:
            sources.append(RagChatSource.model_validate(item))
        except Exception:
            continue

    try:
        trace = json.loads(eval_run.trace_json or "[]")
    except json.JSONDecodeError:
        trace = []

    if not isinstance(trace, list):
        trace = []

    failure_reasons = get_eval_failure_reasons(eval_run)

    return RagEvalRunPublic(
        id=eval_run.id,
        eval_case_id=eval_run.eval_case_id,
        batch_id=eval_run.batch_id,
        knowledge_base_id=eval_run.knowledge_base_id,
        owner_id=eval_run.owner_id,
        question=eval_run.question,
        answer=eval_run.answer,
        expected_keywords=parse_json_list(eval_run.expected_keywords_json),
        expected_source_filename=eval_run.expected_source_filename,
        sources=sources,
        trace=trace,
        keyword_hit=eval_run.keyword_hit,
        source_hit=eval_run.source_hit,
        score=eval_run.score,
        retrieval_type=eval_run.retrieval_type,
        top_k=eval_run.top_k,
        semantic_weight=eval_run.semantic_weight,
        keyword_weight=eval_run.keyword_weight,
        latency_ms=eval_run.latency_ms,
        error_message=eval_run.error_message,
        created_at=eval_run.created_at,
        is_failed=len(failure_reasons) > 0,
        failure_reasons=failure_reasons,
    )


def rag_eval_batch_to_public(batch: RagEvalBatch) -> RagEvalBatchPublic:
    return RagEvalBatchPublic(
        id=batch.id,
        knowledge_base_id=batch.knowledge_base_id,
        owner_id=batch.owner_id,
        name=batch.name,
        note=batch.note,
        top_k=batch.top_k,
        semantic_weight=batch.semantic_weight,
        keyword_weight=batch.keyword_weight,
        total_cases=batch.total_cases,
        ran=batch.ran,
        failed=batch.failed,
        average_score=batch.average_score,
        keyword_hit_rate=batch.keyword_hit_rate,
        source_hit_rate=batch.source_hit_rate,
        average_latency_ms=batch.average_latency_ms,
        created_at=batch.created_at,
    )


def rag_retrieval_preset_to_public(
    preset: RagRetrievalPreset,
) -> RagRetrievalPresetPublic:
    return RagRetrievalPresetPublic(
        id=preset.id,
        knowledge_base_id=preset.knowledge_base_id,
        owner_id=preset.owner_id,
        name=preset.name,
        top_k=preset.top_k,
        semantic_weight=preset.semantic_weight,
        keyword_weight=preset.keyword_weight,
        note=preset.note,
        created_at=preset.created_at,
    )


def model_to_json_dict(item: Any) -> dict[str, Any]:
    if hasattr(item, "model_dump"):
        return item.model_dump(mode="json")

    if hasattr(item, "dict"):
        return item.dict()

    return dict(item)


def agent_tool_calls_to_json(
    tool_calls: list[AgentToolCallPublic],
) -> str:
    return json.dumps(
        [model_to_json_dict(tool_call) for tool_call in tool_calls],
        ensure_ascii=False,
        default=str,
    )


def agent_sources_to_json(
    sources: list[RagChatSource],
) -> str:
    return json.dumps(
        [model_to_json_dict(source) for source in sources],
        ensure_ascii=False,
        default=str,
    )


def agent_trace_to_json(trace: list[str]) -> str:
    return json.dumps(trace, ensure_ascii=False, default=str)


def agent_tool_calls_from_json(
    tool_calls_json: str,
) -> list[AgentToolCallPublic]:
    try:
        raw_items = json.loads(tool_calls_json or "[]")
    except json.JSONDecodeError:
        return []

    return [AgentToolCallPublic(**item) for item in raw_items]


def agent_sources_from_json(
    sources_json: str,
) -> list[RagChatSource]:
    try:
        raw_items = json.loads(sources_json or "[]")
    except json.JSONDecodeError:
        return []

    return [RagChatSource(**item) for item in raw_items]


def agent_trace_from_json(trace_json: str) -> list[str]:
    try:
        raw_items = json.loads(trace_json or "[]")
    except json.JSONDecodeError:
        return []

    return [str(item) for item in raw_items]


def agent_run_to_public(agent_run: AgentRun) -> AgentRunPublic:
    return AgentRunPublic(
        id=agent_run.id,
        knowledge_base_id=agent_run.knowledge_base_id,
        owner_id=agent_run.owner_id,
        repository_analysis_task_id=agent_run.repository_analysis_task_id,
        source_commit_sha=agent_run.source_commit_sha,
        question=agent_run.question,
        answer=agent_run.answer,
        top_k=agent_run.top_k,
        max_steps=agent_run.max_steps,
        semantic_weight=agent_run.semantic_weight,
        keyword_weight=agent_run.keyword_weight,
        latency_ms=agent_run.latency_ms,
        tool_calls=agent_tool_calls_from_json(agent_run.tool_calls_json),
        sources=agent_sources_from_json(agent_run.sources_json),
        trace=agent_trace_from_json(agent_run.trace_json),
        error_message=agent_run.error_message,
        created_at=agent_run.created_at,
    )

def build_agent_settings_public(
    settings: KnowledgeBaseAgentSettings,
) -> KnowledgeBaseAgentSettingsPublic:
    return KnowledgeBaseAgentSettingsPublic(
        id=settings.id,
        knowledge_base_id=settings.knowledge_base_id,
        owner_id=settings.owner_id,
        top_k=settings.top_k,
        max_steps=settings.max_steps,
        semantic_weight=settings.semantic_weight,
        keyword_weight=settings.keyword_weight,
        created_at=settings.created_at,
        updated_at=settings.updated_at,
    )


def get_or_create_agent_settings(
    *,
    session: SessionDep,
    knowledge_base_id: uuid.UUID,
    owner_id: uuid.UUID,
) -> KnowledgeBaseAgentSettings:
    statement = (
        select(KnowledgeBaseAgentSettings)
        .where(KnowledgeBaseAgentSettings.knowledge_base_id == knowledge_base_id)
    )

    settings = session.exec(statement).first()

    if settings:
        return settings

    settings = KnowledgeBaseAgentSettings(
        knowledge_base_id=knowledge_base_id,
        owner_id=owner_id,
        top_k=5,
        max_steps=5,
        semantic_weight=0.75,
        keyword_weight=0.25,
    )

    session.add(settings)
    session.commit()
    session.refresh(settings)

    return settings


def check_keyword_hit(answer: str, expected_keywords: list[str]) -> bool | None:
    keywords = [keyword.strip() for keyword in expected_keywords if keyword.strip()]

    if not keywords:
        return None

    answer_lower = answer.lower()

    return all(keyword.lower() in answer_lower for keyword in keywords)


def check_source_hit(
    *,
    sources: list[RagChatSource],
    expected_source_filename: str | None,
) -> bool | None:
    if not expected_source_filename or not expected_source_filename.strip():
        return None

    expected = expected_source_filename.strip().lower()

    for source in sources:
        filename = source.original_filename.lower()

        if expected in filename or filename in expected:
            return True

    return False


def calculate_eval_score(
    *,
    keyword_hit: bool | None,
    source_hit: bool | None,
) -> float | None:
    checks = []

    if keyword_hit is not None:
        checks.append(keyword_hit)

    if source_hit is not None:
        checks.append(source_hit)

    if not checks:
        return None

    passed = sum(1 for item in checks if item)

    return passed / len(checks)


def calculate_rate(values: list[bool | None]) -> float | None:
    valid_values = [value for value in values if value is not None]

    if not valid_values:
        return None

    passed = sum(1 for value in valid_values if value)

    return passed / len(valid_values)


def calculate_average(values: list[float | int | None]) -> float | None:
    valid_values = [value for value in values if value is not None]

    if not valid_values:
        return None

    return sum(float(value) for value in valid_values) / len(valid_values)


def normalize_retrieval_weights(
    *,
    semantic_weight: float,
    keyword_weight: float,
) -> tuple[float, float]:
    total = semantic_weight + keyword_weight

    if total <= 0:
        raise HTTPException(
            status_code=400,
            detail="semantic_weight and keyword_weight cannot both be 0",
        )

    return semantic_weight / total, keyword_weight / total


def build_eval_summary_from_runs(eval_runs: list[RagEvalRun]) -> RagEvalSummary:
    if not eval_runs:
        return RagEvalSummary(
            total_runs=0,
            average_score=None,
            keyword_hit_rate=None,
            source_hit_rate=None,
            average_latency_ms=None,
            max_score=None,
            min_score=None,
            latest_run_at=None,
        )

    scores = [eval_run.score for eval_run in eval_runs if eval_run.score is not None]

    latest_run = max(
        eval_runs,
        key=lambda eval_run: eval_run.created_at
    )

    return RagEvalSummary(
        total_runs=len(eval_runs),
        average_score=calculate_average([eval_run.score for eval_run in eval_runs]),
        keyword_hit_rate=calculate_rate(
            [eval_run.keyword_hit for eval_run in eval_runs]
        ),
        source_hit_rate=calculate_rate(
            [eval_run.source_hit for eval_run in eval_runs]
        ),
        average_latency_ms=calculate_average(
            [eval_run.latency_ms for eval_run in eval_runs]
        ),
        max_score=max(scores) if scores else None,
        min_score=min(scores) if scores else None,
        latest_run_at=latest_run.created_at,
    )


def run_rag_once_for_compare(
    *,
    session: SessionDep,
    knowledge_base_id: uuid.UUID,
    question: str,
    top_k: int,
    semantic_weight: float,
    keyword_weight: float,
) -> tuple[str, list[RagChatSource], list[str], int]:
    started_at = time.perf_counter()
    trace: list[str] = []

    trace.append(
        f"compare_rag_params_top_k={top_k},semantic_weight={semantic_weight:.2f},keyword_weight={keyword_weight:.2f}"
    )

    hybrid_rows = hybrid_retrieve_chunks_for_rag(
        session=session,
        knowledge_base_id=knowledge_base_id,
        query=question,
        top_k=top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
    )

    trace.append("hybrid_retrieve_chunks")

    sources = [
        RagChatSource(
            document_id=document.id,
            original_filename=document.original_filename,
            chunk_id=chunk.id,
            chunk_index=chunk.chunk_index,
            content=chunk.content,
            content_length=chunk.content_length,
            match_count=match_count,
            similarity=similarity,
            retrieval_type="hybrid",
        )
        for chunk, document, similarity, match_count, hybrid_score in hybrid_rows
    ]

    trace.append("build_hybrid_sources")

    answer = generate_rag_answer(
        question=question,
        sources=sources,
    )

    trace.append("generate_rag_answer")

    latency_ms = int((time.perf_counter() - started_at) * 1000)

    return answer, sources, trace, latency_ms


def run_agent_once_for_compare(
    *,
    session: SessionDep,
    knowledge_base_id: uuid.UUID,
    question: str,
    top_k: int,
    max_steps: int,
    semantic_weight: float,
    keyword_weight: float,
) -> tuple[
    str,
    list[AgentToolCallPublic],
    list[RagChatSource],
    list[str],
    int,
]:
    started_at = time.perf_counter()

    answer, tool_calls, sources, trace = (
        run_knowledge_base_agent(
            session=session,
            knowledge_base_id=
            knowledge_base_id,
            repository_analysis_task_id=
            repository_scope_id,
            question=question,
            top_k=request.top_k,
            max_steps=request.max_steps,
            semantic_weight=
            request.semantic_weight,
            keyword_weight=
            request.keyword_weight,
        )
    )

    latency_ms = int((time.perf_counter() - started_at) * 1000)

    return answer, tool_calls, sources, trace, latency_ms


def build_rag_agent_compare_summary(
    items: list[RagAgentCompareEvalItem],
) -> RagAgentCompareEvalSummary:
    if not items:
        return RagAgentCompareEvalSummary()

    all_tool_names: list[str] = []

    for item in items:
        all_tool_names.extend(item.agent_tool_names)

    return RagAgentCompareEvalSummary(
        total_cases=len(items),
        ran=len(items),
        failed=sum(1 for item in items if item.is_failed),
        average_rag_latency_ms=calculate_average(
            [item.rag_latency_ms for item in items]
        ),
        average_agent_latency_ms=calculate_average(
            [item.agent_latency_ms for item in items]
        ),
        average_rag_sources_count=calculate_average(
            [item.rag_sources_count for item in items]
        ),
        average_agent_sources_count=calculate_average(
            [item.agent_sources_count for item in items]
        ),
        average_agent_tool_call_count=calculate_average(
            [item.agent_tool_call_count for item in items]
        ),
        total_agent_failed_tool_count=sum(
            item.agent_failed_tool_count for item in items
        ),
        summarize_document_count=sum(
            1 for name in all_tool_names if name == "summarize_document"
        ),
        read_document_chunks_count=sum(
            1 for name in all_tool_names if name == "read_document_chunks"
        ),
        search_rag_history_count=sum(
            1 for name in all_tool_names if name == "search_rag_history"
        ),
    )


def save_rag_agent_compare_batch(
    *,
    session: SessionDep,
    knowledge_base_id: uuid.UUID,
    owner_id: uuid.UUID,
    request: RagAgentCompareEvalRequest,
    summary: RagAgentCompareEvalSummary,
    items: list[RagAgentCompareEvalItem],
) -> RagAgentCompareBatch:
    batch_name = request.name or "RAG vs Agent Compare"

    batch = RagAgentCompareBatch(
        knowledge_base_id=knowledge_base_id,
        owner_id=owner_id,
        name=batch_name,
        top_k=request.top_k,
        max_steps=request.max_steps,
        limit=request.limit,
        semantic_weight=request.semantic_weight,
        keyword_weight=request.keyword_weight,
        total_cases=summary.total_cases,
        ran=summary.ran,
        failed=summary.failed,
        average_rag_latency_ms=summary.average_rag_latency_ms,
        average_agent_latency_ms=summary.average_agent_latency_ms,
        average_rag_sources_count=summary.average_rag_sources_count,
        average_agent_sources_count=summary.average_agent_sources_count,
        average_agent_tool_call_count=summary.average_agent_tool_call_count,
        total_agent_failed_tool_count=summary.total_agent_failed_tool_count,
        summarize_document_count=summary.summarize_document_count,
        read_document_chunks_count=summary.read_document_chunks_count,
        search_rag_history_count=summary.search_rag_history_count,
    )

    session.add(batch)
    session.flush()

    for item in items:
        db_item = RagAgentCompareItem(
            batch_id=batch.id,
            knowledge_base_id=knowledge_base_id,
            owner_id=owner_id,
            eval_case_id=item.eval_case_id,
            question=item.question,
            expected_keywords_json=dump_json_data(item.expected_keywords),
            expected_source_filename=item.expected_source_filename,
            rag_answer=item.rag_answer,
            agent_answer=item.agent_answer,
            rag_latency_ms=item.rag_latency_ms,
            agent_latency_ms=item.agent_latency_ms,
            rag_sources_count=item.rag_sources_count,
            agent_sources_count=item.agent_sources_count,
            agent_tool_call_count=item.agent_tool_call_count,
            agent_failed_tool_count=item.agent_failed_tool_count,
            agent_tool_names_json=dump_json_data(item.agent_tool_names),
            rag_sources_json=dump_model_list(item.rag_sources),
            agent_sources_json=dump_model_list(item.agent_sources),
            agent_tool_calls_json=dump_model_list(item.agent_tool_calls),
            rag_trace_json=dump_json_data(item.rag_trace),
            agent_trace_json=dump_json_data(item.agent_trace),
            rag_error_message=item.rag_error_message,
            agent_error_message=item.agent_error_message,
            is_failed=item.is_failed,
        )

        session.add(db_item)

    session.commit()
    session.refresh(batch)

    return batch


def get_eval_failure_reasons(eval_run: RagEvalRun) -> list[str]:
    reasons: list[str] = []

    if eval_run.error_message:
        reasons.append("运行错误")

    if eval_run.keyword_hit is False:
        reasons.append("关键词未命中")

    if eval_run.source_hit is False:
        reasons.append("来源未命中")

    if eval_run.score is not None and eval_run.score < 1:
        reasons.append("得分未满")

    if eval_run.score == 0:
        reasons.append("零分样例")

    return reasons


def is_eval_run_failed(eval_run: RagEvalRun) -> bool:
    return len(get_eval_failure_reasons(eval_run)) > 0


def build_retrieval_param_key(
    *,
    top_k: int,
    semantic_weight: float,
    keyword_weight: float,
) -> str:
    return (
        f"{top_k}:"
        f"{round(float(semantic_weight), 4)}:"
        f"{round(float(keyword_weight), 4)}"
    )


def build_eval_param_groups_from_runs(
    *,
    eval_runs: list[RagEvalRun],
    presets: list[RagRetrievalPreset],
) -> list[RagEvalParamGroup]:
    groups: dict[str, list[RagEvalRun]] = {}

    preset_names_by_key: dict[str, list[str]] = {}

    for preset in presets:
        preset_key = build_retrieval_param_key(
            top_k=preset.top_k,
            semantic_weight=preset.semantic_weight,
            keyword_weight=preset.keyword_weight,
        )

        if preset_key not in preset_names_by_key:
            preset_names_by_key[preset_key] = []

        preset_names_by_key[preset_key].append(preset.name)

    for eval_run in eval_runs:
        group_key = build_retrieval_param_key(
            top_k=eval_run.top_k,
            semantic_weight=eval_run.semantic_weight,
            keyword_weight=eval_run.keyword_weight,
        )

        if group_key not in groups:
            groups[group_key] = []

        groups[group_key].append(eval_run)

    results: list[RagEvalParamGroup] = []

    for group_key, runs in groups.items():
        first_run = runs[0]

        failed_count = sum(1 for run in runs if is_eval_run_failed(run))
        failure_rate = failed_count / len(runs) if runs else None

        latest_run = runs[0]

        for run in runs:
            if run.created_at and latest_run.created_at:
                if run.created_at > latest_run.created_at:
                    latest_run = run

        results.append(
            RagEvalParamGroup(
                top_k=first_run.top_k,
                semantic_weight=round(float(first_run.semantic_weight), 4),
                keyword_weight=round(float(first_run.keyword_weight), 4),
                preset_names=preset_names_by_key.get(group_key, []),
                total_runs=len(runs),
                average_score=calculate_average([run.score for run in runs]),
                keyword_hit_rate=calculate_rate(
                    [run.keyword_hit for run in runs]
                ),
                source_hit_rate=calculate_rate(
                    [run.source_hit for run in runs]
                ),
                failure_rate=failure_rate,
                average_latency_ms=calculate_average(
                    [run.latency_ms for run in runs]
                ),
                latest_run_at=latest_run.created_at,
            )
        )

    results.sort(
        key=lambda item: (
            -(item.average_score or 0),
            item.failure_rate if item.failure_rate is not None else 1,
            item.average_latency_ms if item.average_latency_ms is not None else 999999,
        )
    )

    return results


def normalize_zip_member_path(raw_path: str) -> str:
    normalized = raw_path.replace("\\", "/").strip()

    while normalized.startswith("./"):
        normalized = normalized[2:]

    return normalized


def is_safe_zip_member_path(path: str) -> bool:
    if not path:
        return False

    if path.startswith("/"):
        return False

    path_parts = Path(path).parts

    if ".." in path_parts:
        return False

    return True


def should_ignore_repository_path(path: str) -> bool:
    normalized_path = normalize_zip_member_path(path)

    if not is_safe_zip_member_path(normalized_path):
        return True

    path_obj = Path(normalized_path)
    parts = set(path_obj.parts)

    if parts & REPOSITORY_IGNORE_DIRS:
        return True

    filename = path_obj.name

    if filename in REPOSITORY_IGNORE_FILE_PREFIXES:
        return True

    suffix = path_obj.suffix.lower()

    if suffix in REPOSITORY_IGNORE_FILE_SUFFIXES:
        return True

    return False


def is_supported_repository_code_file(path: str) -> bool:
    normalized_path = normalize_zip_member_path(path)

    if should_ignore_repository_path(normalized_path):
        return False

    return is_code_file(normalized_path)


def build_repository_original_filename(
    *,
    repository_name: str,
    member_path: str,
) -> str:
    safe_repository_name = repository_name.strip() or "repository"
    safe_repository_name = safe_repository_name.replace("\\", "/").split("/")[-1]

    normalized_member_path = normalize_zip_member_path(member_path)

    return f"{safe_repository_name}/{normalized_member_path}"


def validate_file_type(file: UploadFile) -> str:
    original_filename = file.filename or "uploaded_file"
    suffix = Path(original_filename).suffix.lower()

    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                f"Supported extensions: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
            ),
        )

    if suffix in CODE_EXTENSIONS:
        return suffix

    allowed_content_types = ALLOWED_CONTENT_TYPES.get(suffix, set())

    if file.content_type and file.content_type not in allowed_content_types:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file content type",
        )

    return suffix


def decode_text_content(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue

    return content.decode("utf-8", errors="ignore")


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)

    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")

    return text.strip()


def extract_text_from_file(
    *,
    original_filename: str,
    content: bytes,
) -> str | None:
    suffix = Path(original_filename).suffix.lower()

    if suffix not in TEXT_EXTENSIONS:
        return None

    return decode_text_content(content)


def split_text_into_chunks(
    text: str,
    *,
    chunk_size: int = CHUNK_SIZE,
    chunk_overlap: int = CHUNK_OVERLAP,
) -> list[str]:
    text = normalize_text(text)

    if not text:
        return []

    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    chunks: list[str] = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)

        if end < text_length:
            split_at = text.rfind("\n\n", start, end)

            if split_at == -1 or split_at <= start + chunk_size // 2:
                split_at = text.rfind("\n", start, end)

            if split_at != -1 and split_at > start:
                end = split_at

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= text_length:
            break

        next_start = end - chunk_overlap

        if next_start <= start:
            next_start = end

        start = next_start

    return chunks

def extract_code_chunk_metadata_for_eval(content: str) -> dict[str, str]:
    metadata: dict[str, str] = {}

    for line in content.splitlines()[:8]:
        if ":" not in line:
            continue

        key, value = line.split(":", 1)
        key = key.strip()
        value = value.strip()

        if key in {
            "file_path",
            "language",
            "symbol_name",
            "symbol_type",
            "line_range",
        }:
            metadata[key] = value

    return metadata


def is_code_document_filename(filename: str) -> bool:
    suffix = Path(filename).suffix.lower()

    return suffix in CODE_EXTENSIONS


def build_code_eval_cases_from_documents(
    *,
    session: SessionDep,
    knowledge_base_id: uuid.UUID,
    owner_id: uuid.UUID,
    max_files: int = 20,
    max_symbols: int = 40,
) -> list[RagEvalCase]:
    statement = (
        select(Document)
        .where(Document.knowledge_base_id == knowledge_base_id)
        .order_by(col(Document.original_filename))
        .limit(max_files * 3)
    )

    documents = session.exec(statement).all()

    code_documents = [
        document
        for document in documents
        if is_code_document_filename(document.original_filename)
    ]

    code_documents = code_documents[:max_files]

    created_cases: list[RagEvalCase] = []

    if code_documents:
        overview_case = RagEvalCase(
            knowledge_base_id=knowledge_base_id,
            owner_id=owner_id,
            question="这个代码仓库中有哪些主要代码文件？请按文件路径简要说明。",
            expected_keywords_json=json.dumps(
                [
                    Path(document.original_filename).name
                    for document in code_documents[:5]
                ],
                ensure_ascii=False,
            ),
            expected_source_filename=code_documents[0].original_filename,
            note="code:file_overview",
        )

        session.add(overview_case)
        created_cases.append(overview_case)

    symbol_count = 0

    for document in code_documents:
        file_question = RagEvalCase(
            knowledge_base_id=knowledge_base_id,
            owner_id=owner_id,
            question=f"{document.original_filename} 这个代码文件的主要作用是什么？",
            expected_keywords_json=json.dumps(
                [
                    Path(document.original_filename).name,
                ],
                ensure_ascii=False,
            ),
            expected_source_filename=document.original_filename,
            note="code:file_overview",
        )

        session.add(file_question)
        created_cases.append(file_question)

        chunk_statement = (
            select(DocumentChunk)
            .where(DocumentChunk.document_id == document.id)
            .order_by(col(DocumentChunk.chunk_index))
            .limit(50)
        )

        chunks = session.exec(chunk_statement).all()

        for chunk in chunks:
            metadata = extract_code_chunk_metadata_for_eval(chunk.content)

            symbol_name = metadata.get("symbol_name")
            symbol_type = metadata.get("symbol_type")
            line_range = metadata.get("line_range")

            if not symbol_name or symbol_name == "-":
                continue

            if not symbol_type or symbol_type == "-":
                continue

            if symbol_count >= max_symbols:
                break

            symbol_question = RagEvalCase(
                knowledge_base_id=knowledge_base_id,
                owner_id=owner_id,
                question=(
                    f"{symbol_name} 这个 {symbol_type} 在哪里定义？"
                    f"它的主要作用是什么？"
                ),
                expected_keywords_json=json.dumps(
                    [
                        symbol_name,
                        symbol_type,
                        line_range or "",
                    ],
                    ensure_ascii=False,
                ),
                expected_source_filename=document.original_filename,
                note="code:symbol_definition",
            )

            session.add(symbol_question)
            created_cases.append(symbol_question)

            reference_question = RagEvalCase(
                knowledge_base_id=knowledge_base_id,
                owner_id=owner_id,
                question=f"{symbol_name} 在代码仓库中出现在哪里？请说明可能的定义或引用位置。",
                expected_keywords_json=json.dumps(
                    [
                        symbol_name,
                        Path(document.original_filename).name,
                    ],
                    ensure_ascii=False,
                ),
                expected_source_filename=document.original_filename,
                note="code:symbol_reference",
            )

            session.add(reference_question)
            created_cases.append(reference_question)

            logic_question = RagEvalCase(
                knowledge_base_id=knowledge_base_id,
                owner_id=owner_id,
                question=(
                    f"请结合所在文件上下文，解释 {symbol_name} 这个 {symbol_type} "
                    f"的完整执行逻辑或主要实现流程。"
                ),
                expected_keywords_json=json.dumps(
                    [
                        symbol_name,
                        symbol_type,
                        Path(document.original_filename).name,
                    ],
                    ensure_ascii=False,
                ),
                expected_source_filename=document.original_filename,
                note="code:code_logic",
            )

            session.add(logic_question)
            created_cases.append(logic_question)

            symbol_count += 1

        if symbol_count >= max_symbols:
            break

    return created_cases


def create_document_chunks(
    *,
    session: SessionDep,
    document: Document,
    content: bytes,
) -> int:
    if is_code_file(document.original_filename):
        code_chunks = split_code_into_chunks(
            filename=document.original_filename,
            file_bytes=content,
        )

        for index, code_chunk in enumerate(code_chunks):
            document_chunk = DocumentChunk(
                document_id=document.id,
                knowledge_base_id=document.knowledge_base_id,
                owner_id=document.owner_id,
                chunk_index=index,
                content=code_chunk.content,
                content_length=len(code_chunk.content),
            )
            session.add(document_chunk)

            embed_document_chunk(
                session=session,
                chunk=document_chunk,
            )

        return len(code_chunks)

    extracted_text = parse_document_content(
        filename=document.original_filename,
        file_bytes=content,
    )

    chunks = split_text_into_chunks(extracted_text)

    for index, chunk_content in enumerate(chunks):
        document_chunk = DocumentChunk(
            document_id=document.id,
            knowledge_base_id=document.knowledge_base_id,
            owner_id=document.owner_id,
            chunk_index=index,
            content=chunk_content,
            content_length=len(chunk_content),
        )
        session.add(document_chunk)

        embed_document_chunk(
            session=session,
            chunk=document_chunk,
        )

    return len(chunks)

@router.get("/documents/{document_id}/download")
def download_document(
    session: SessionDep,
    current_user: CurrentUser,
    document_id: uuid.UUID,
) -> FileResponse:
    document = get_document_or_404(
        session=session,
        document_id=document_id,
    )
    check_document_permission(
        document=document,
        current_user=current_user,
    )

    storage_path = Path(document.storage_path)

    if not storage_path.exists() or not storage_path.is_file():
        raise HTTPException(status_code=404, detail="File not found")

    return FileResponse(
        path=storage_path,
        filename=document.original_filename,
        media_type=document.content_type or "application/octet-stream",
    )


@router.post(
    "/knowledge-bases/{knowledge_base_id}/documents/",
    response_model=DocumentPublic,
)
async def upload_document(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    file: UploadFile = File(...),
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    suffix = validate_file_type(file)

    content = await file.read()
    file_size = len(content)

    if file_size == 0:
        raise HTTPException(status_code=400, detail="Empty file is not allowed")

    if file_size > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File too large")

    document_id = uuid.uuid4()
    filename = f"{document_id}{suffix}"

    storage_dir = STORAGE_ROOT / str(knowledge_base.owner_id) / str(knowledge_base.id)
    storage_dir.mkdir(parents=True, exist_ok=True)

    storage_path = storage_dir / filename

    try:
        storage_path.write_bytes(content)
    except OSError:
        raise HTTPException(status_code=500, detail="Could not save uploaded file")

    document = Document(
        id=document_id,
        knowledge_base_id=knowledge_base.id,
        owner_id=knowledge_base.owner_id,
        filename=filename,
        original_filename=file.filename or filename,
        content_type=file.content_type,
        file_size=file_size,
        storage_path=str(storage_path),
        status="uploaded",
    )

    session.add(document)

    try:
        chunk_count = create_document_chunks(
            session=session,
            document=document,
            content=content,
        )
    except DocumentParseError as error:
        chunk_count = 0
        document.status = "uploaded"
        document.error_message = str(error)

    if chunk_count > 0:
        document.status = "parsed"
        document.error_message = None
    elif not document.error_message:
        document.status = "uploaded"
        document.error_message = "No extractable text found"

    session.add(document)
    session.commit()
    session.refresh(document)

    return document


SUPPORTED_REPOSITORY_ARCHIVE_SUFFIXES = {
    ".zip",
    ".rar",
}


def is_supported_repository_archive_file(filename: str) -> bool:
    return Path(filename.lower()).suffix in SUPPORTED_REPOSITORY_ARCHIVE_SUFFIXES


def rar_member_is_dir(member: rarfile.RarInfo) -> bool:
    isdir = getattr(member, "isdir", None)

    if callable(isdir):
        return bool(isdir())

    return member.filename.endswith("/")


def get_archive_member_filename(member: Any) -> str:
    return str(getattr(member, "filename", ""))


def get_archive_member_file_size(member: Any) -> int:
    if hasattr(member, "file_size"):
        return int(getattr(member, "file_size") or 0)

    if hasattr(member, "file_size"):
        return int(getattr(member, "file_size") or 0)

    return int(getattr(member, "file_size", 0) or 0)


def archive_member_is_dir(
    *,
    archive_suffix: str,
    member: Any,
) -> bool:
    if archive_suffix == ".zip":
        return bool(member.is_dir())

    if archive_suffix == ".rar":
        return rar_member_is_dir(member)

    return False


def read_archive_member_bytes(
    *,
    archive_file: Any,
    archive_suffix: str,
    member: Any,
) -> bytes:
    if archive_suffix == ".zip":
        return archive_file.read(member)

    if archive_suffix == ".rar":
        with archive_file.open(member) as member_file:
            return member_file.read()

    raise HTTPException(
        status_code=400,
        detail="Unsupported repository archive type",
    )

async def save_upload_file_to_path(
    *,
    upload_file: UploadFile,
    target_path: Path,
    max_size: int,
) -> int:
    total_size = 0

    with target_path.open("wb") as buffer:
        while True:
            chunk = await upload_file.read(UPLOAD_COPY_CHUNK_SIZE)

            if not chunk:
                break

            total_size += len(chunk)

            if total_size > max_size:
                raise HTTPException(
                    status_code=400,
                    detail="Repository archive file too large",
                )

            buffer.write(chunk)

    return total_size




@router.post(
    "/knowledge-bases/{knowledge_base_id}/code-repository-zip",
    response_model=DocumentsPublic,
    operation_id="upload_code_repository_zip",
)

async def upload_code_repository_zip(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    file: UploadFile = File(...),
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    original_archive_filename = file.filename or "repository.zip"
    archive_suffix = Path(original_archive_filename).suffix.lower()

    if not is_supported_repository_archive_file(original_archive_filename):
        raise HTTPException(
            status_code=400,
            detail="Only .zip and .rar repository files are supported",
        )

    repository_name = Path(original_archive_filename).stem or "repository"

    created_documents: list[Document] = []
    total_code_size = 0

    storage_dir = STORAGE_ROOT / str(knowledge_base.owner_id) / str(knowledge_base.id)
    storage_dir.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as temporary_dir:
        temporary_path = Path(temporary_dir)
        archive_path = temporary_path / f"{uuid.uuid4()}{archive_suffix}"

        archive_size = await save_upload_file_to_path(
            upload_file=file,
            target_path=archive_path,
            max_size=MAX_REPOSITORY_ARCHIVE_SIZE,
        )

        if archive_size == 0:
            raise HTTPException(
                status_code=400,
                detail="Empty repository archive file is not allowed",
            )

        try:
            if archive_suffix == ".zip":
                archive_file = zipfile.ZipFile(archive_path)
            elif archive_suffix == ".rar":
                archive_file = rarfile.RarFile(archive_path)
            else:
                raise HTTPException(
                    status_code=400,
                    detail="Unsupported repository archive type",
                )
        except zipfile.BadZipFile as error:
            raise HTTPException(
                status_code=400,
                detail=f"Invalid zip file: {error}",
            ) from error
        except rarfile.RarCannotExec as error:
            raise HTTPException(
                status_code=500,
                detail=(
                    "RAR extraction tool is not available in backend container. "
                    "Please install unar or unrar."
                ),
            ) from error
        except rarfile.PasswordRequired as error:
            raise HTTPException(
                status_code=400,
                detail="Password-protected rar files are not supported",
            ) from error
        except rarfile.BadRarFile as error:
            raise HTTPException(
                status_code=400,
                detail=f"Invalid rar file: {error}",
            ) from error
        except rarfile.Error as error:
            raise HTTPException(
                status_code=400,
                detail=f"Could not open rar file: {error}",
            ) from error

        with archive_file:
            code_members = []

            for member in archive_file.infolist():
                if archive_member_is_dir(
                    archive_suffix=archive_suffix,
                    member=member,
                ):
                    continue

                member_filename = get_archive_member_filename(member)
                member_path = normalize_zip_member_path(member_filename)

                if not is_supported_repository_code_file(member_path):
                    continue

                member_file_size = get_archive_member_file_size(member)

                if member_file_size <= 0:
                    continue

                if member_file_size > MAX_REPOSITORY_SINGLE_CODE_FILE_SIZE:
                    continue

                total_code_size += member_file_size

                if total_code_size > MAX_REPOSITORY_TOTAL_CODE_SIZE:
                    raise HTTPException(
                        status_code=400,
                        detail="Repository code files are too large",
                    )

                code_members.append(member)

                if len(code_members) >= MAX_REPOSITORY_FILE_COUNT:
                    break

            if not code_members:
                raise HTTPException(
                    status_code=400,
                    detail="No supported code files found in repository archive",
                )

            for member in code_members:
                member_filename = get_archive_member_filename(member)
                member_path = normalize_zip_member_path(member_filename)

                try:
                    file_bytes = read_archive_member_bytes(
                        archive_file=archive_file,
                        archive_suffix=archive_suffix,
                        member=member,
                    )
                except rarfile.RarCannotExec as error:
                    raise HTTPException(
                        status_code=500,
                        detail=(
                            "RAR extraction tool is not available in backend container. "
                            "Please install unar or unrar."
                        ),
                    ) from error
                except rarfile.PasswordRequired as error:
                    raise HTTPException(
                        status_code=400,
                        detail="Password-protected rar files are not supported",
                    ) from error
                except rarfile.Error:
                    continue
                except Exception:
                    continue

                if not file_bytes:
                    continue

                if len(file_bytes) > MAX_REPOSITORY_SINGLE_CODE_FILE_SIZE:
                    continue

                original_filename = build_repository_original_filename(
                    repository_name=repository_name,
                    member_path=member_path,
                )

                suffix = Path(original_filename).suffix.lower()
                document_id = uuid.uuid4()
                stored_filename = f"{document_id}{suffix}"
                storage_path = storage_dir / stored_filename

                try:
                    storage_path.write_bytes(file_bytes)
                except OSError:
                    raise HTTPException(
                        status_code=500,
                        detail=f"Could not save file from repository: {original_filename}",
                    )

                document = Document(
                    id=document_id,
                    knowledge_base_id=knowledge_base.id,
                    owner_id=knowledge_base.owner_id,
                    filename=stored_filename,
                    original_filename=original_filename,
                    content_type="text/plain",
                    file_size=len(file_bytes),
                    storage_path=str(storage_path),
                    status="uploaded",
                )

                session.add(document)

                try:
                    chunk_count = create_document_chunks(
                        session=session,
                        document=document,
                        content=file_bytes,
                    )
                except DocumentParseError as error:
                    chunk_count = 0
                    document.status = "uploaded"
                    document.error_message = str(error)

                if chunk_count > 0:
                    document.status = "parsed"
                    document.error_message = None
                elif not document.error_message:
                    document.status = "uploaded"
                    document.error_message = "No extractable code found"

                session.add(document)
                created_documents.append(document)

    session.commit()

    for document in created_documents:
        session.refresh(document)

    return DocumentsPublic(
        data=[
            DocumentPublic.model_validate(document)
            for document in created_documents
        ],
        count=len(created_documents),
    )



@router.get(
    "/knowledge-bases/{knowledge_base_id}/documents/",
    response_model=DocumentsPublic,
)
def read_documents_by_knowledge_base(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 100,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    count_statement = (
        select(func.count())
        .select_from(Document)
        .where(Document.knowledge_base_id == knowledge_base_id)
    )
    count = session.exec(count_statement).one()

    statement = (
        select(Document)
        .where(Document.knowledge_base_id == knowledge_base_id)
        .order_by(col(Document.created_at).desc())
        .offset(skip)
        .limit(limit)
    )
    documents = session.exec(statement).all()

    documents_public = [
        DocumentPublic.model_validate(document)
        for document in documents
    ]

    return DocumentsPublic(data=documents_public, count=count)


@router.get(
    "/documents/{document_id}/chunks",
    response_model=DocumentChunksPublic,
)
def read_document_chunks(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    document_id: uuid.UUID,
    skip: int = 0,
    limit: int = 100,
) -> Any:
    document = get_document_or_404(
        session=session,
        document_id=document_id,
    )
    check_document_permission(
        document=document,
        current_user=current_user,
    )

    count_statement = (
        select(func.count())
        .select_from(DocumentChunk)
        .where(DocumentChunk.document_id == document_id)
    )
    count = session.exec(count_statement).one()

    statement = (
        select(DocumentChunk)
        .where(DocumentChunk.document_id == document_id)
        .order_by(col(DocumentChunk.chunk_index))
        .offset(skip)
        .limit(limit)
    )
    chunks = session.exec(statement).all()

    chunks_public = [
        DocumentChunkPublic.model_validate(chunk)
        for chunk in chunks
    ]

    return DocumentChunksPublic(data=chunks_public, count=count)


@router.post(
    "/knowledge-bases/{knowledge_base_id}/search",
    response_model=KnowledgeBaseSearchResults,
)
def search_knowledge_base_chunks(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: KnowledgeBaseSearchRequest,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    query = request.query.strip()

    if not query:
        raise HTTPException(status_code=400, detail="Search query is required")

    terms = build_search_terms(query)

    search_conditions = []

    for term in terms:
        pattern = f"%{term}%"
        search_conditions.append(col(DocumentChunk.content).ilike(pattern))
        search_conditions.append(col(Document.original_filename).ilike(pattern))

    count_statement = (
        select(func.count())
        .select_from(DocumentChunk)
        .join(Document, Document.id == DocumentChunk.document_id)
        .where(DocumentChunk.knowledge_base_id == knowledge_base_id)
        .where(or_(*search_conditions))
    )
    count = session.exec(count_statement).one()

    statement = (
        select(DocumentChunk, Document)
        .join(Document, Document.id == DocumentChunk.document_id)
        .where(DocumentChunk.knowledge_base_id == knowledge_base_id)
        .where(or_(*search_conditions))
        .order_by(col(DocumentChunk.chunk_index))
        .limit(request.limit * 3)
    )

    rows = session.exec(statement).all()

    results: list[KnowledgeBaseSearchResult] = []

    for chunk, document in rows:
        match_count = count_keyword_matches(chunk.content, terms)

        results.append(
            KnowledgeBaseSearchResult(
                document_id=document.id,
                original_filename=document.original_filename,
                chunk_id=chunk.id,
                chunk_index=chunk.chunk_index,
                content=chunk.content,
                content_length=chunk.content_length,
                match_count=match_count,
            )
        )

    results.sort(
        key=lambda result: (
            -result.match_count,
            result.original_filename,
            result.chunk_index,
        )
    )

    return KnowledgeBaseSearchResults(
        data=results[: request.limit],
        count=count,
    )

@router.post(
    "/knowledge-bases/{knowledge_base_id}/semantic-search",
    response_model=KnowledgeBaseSemanticSearchResults,
    operation_id="semantic_search_knowledge_base_chunks",
)
def semantic_search_knowledge_base_chunks(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: KnowledgeBaseSemanticSearchRequest,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    query = request.query.strip()

    if not query:
        raise HTTPException(status_code=400, detail="Search query is required")

    try:
        rows = semantic_search_chunks(
            session=session,
            knowledge_base_id=knowledge_base_id,
            query=query,
            top_k=request.top_k,
        )
    except EmbeddingError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error

    results = [
        KnowledgeBaseSemanticSearchResult(
            document_id=document.id,
            original_filename=document.original_filename,
            chunk_id=chunk.id,
            chunk_index=chunk.chunk_index,
            content=chunk.content,
            content_length=chunk.content_length,
            similarity=similarity,
        )
        for chunk, document, similarity in rows
    ]

    return KnowledgeBaseSemanticSearchResults(
        data=results,
        count=len(results),
    )


@router.post(
    "/knowledge-bases/{knowledge_base_id}/chat",
    response_model=RagChatResponse,
    operation_id="chat_with_knowledge_base",
)
def chat_with_knowledge_base(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: RagChatRequest,
) -> Any:
    started_at = time.perf_counter()
    trace: list[str] = []

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    question = request.question.strip()

    if not question:
        raise HTTPException(status_code=400, detail="Question is required")

    repository_scope_id = (
        resolve_repository_analysis_scope(
            session=session,
            current_user=current_user,
            knowledge_base_id=
            knowledge_base_id,
            repository_analysis_task_id=
            request
                .repository_analysis_task_id,
        )
    )
    repository_scope_task = None

    if repository_scope_id is not None:
        repository_scope_task = (
            session.get(
                RepositoryAnalysisTask,
                repository_scope_id,
            )
        )

        if repository_scope_task is None:
            raise HTTPException(
                status_code=404,
                detail={
                    "code": (
                        "REPOSITORY_ANALYSIS_"
                        "NOT_FOUND"
                    ),
                    "message": (
                        "Repository analysis "
                        "task not found"
                    ),
                },
            )
    if repository_scope_id is not None:
        trace.append(
            "repository_scope="
            f"{repository_scope_id}",
        )
    trace.append("check_permission")
    semantic_weight, keyword_weight = normalize_retrieval_weights(
        semantic_weight=request.semantic_weight,
        keyword_weight=request.keyword_weight,
    )

    trace.append(
        f"retrieval_params_top_k={request.top_k},semantic_weight={semantic_weight:.2f},keyword_weight={keyword_weight:.2f}"
    )

    hybrid_rows = (
        hybrid_retrieve_chunks_for_rag(
            session=session,
            knowledge_base_id=(
                knowledge_base_id
            ),
            query=question,
            top_k=request.top_k,
            semantic_weight=(
                semantic_weight
            ),
            keyword_weight=(
                keyword_weight
            ),
            repository_analysis_task_id=(
                repository_scope_id
            ),
        )
    )

    trace.append("hybrid_retrieve_chunks")

    sources = [
        build_rag_chat_source(
            chunk=chunk,
            document=document,
            match_count=match_count,
            similarity=similarity,
            retrieval_type="hybrid",
        )
        for (
            chunk,
            document,
            similarity,
            match_count,
            _hybrid_score,
        ) in hybrid_rows
    ]

    trace.append("build_hybrid_sources")

    answer = generate_rag_answer(
        question=question,
        sources=sources,
    )

    trace.append("generate_answer")

    latency_ms = int((time.perf_counter() - started_at) * 1000)

    retrieval_type = "none"

    if sources:
        retrieval_type = sources[0].retrieval_type

    rag_run = RagRun(
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
        repository_analysis_task_id=(
            repository_scope_id
        ),

        source_commit_sha=(
            repository_scope_task
                .resolved_commit_sha
            if repository_scope_task
               is not None
            else None
        ),
        question=question,
        answer=answer,
        retrieval_type=retrieval_type,
        top_k=request.top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
        latency_ms=latency_ms,
        sources_json=json.dumps(
            jsonable_encoder(sources),
            ensure_ascii=False,
        ),
        trace_json=json.dumps(
            trace,
            ensure_ascii=False,
        ),
    )

    session.add(rag_run)
    session.commit()
    session.refresh(rag_run)

    return RagChatResponse(
        answer=answer,
        sources=sources,
        trace=trace,
        run_id=rag_run.id,
        latency_ms=latency_ms,
    )


@router.post(
    "/knowledge-bases/{knowledge_base_id}/agent-chat",
    response_model=AgentChatResponse,
    operation_id="agent_chat_with_knowledge_base",
)
def agent_chat_with_knowledge_base(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: AgentChatRequest,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    repository_scope_id = (
        resolve_repository_analysis_scope(
            session=session,
            current_user=current_user,
            knowledge_base_id=
            knowledge_base_id,
            repository_analysis_task_id=
            request
                .repository_analysis_task_id,
        )
    )
    repository_scope_task = None

    if repository_scope_id is not None:
        repository_scope_task = (
            session.get(
                RepositoryAnalysisTask,
                repository_scope_id,
            )
        )

        if repository_scope_task is None:
            raise HTTPException(
                status_code=404,
                detail={
                    "code": (
                        "REPOSITORY_ANALYSIS_"
                        "NOT_FOUND"
                    ),
                    "message": (
                        "Repository analysis "
                        "task not found"
                    ),
                },
            )

    question = request.question.strip()

    if not question:
        raise HTTPException(status_code=400, detail="Question is required")

    if (
            repository_scope_id
            is not None
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code":
                    "REPOSITORY_AGENT_SCOPE_NOT_READY",
                "message": (
                    "Repository-scoped Agent is not enabled yet. "
                    "Agent tools must enforce repository task scope "
                    "before this request can run."
                ),
            },
        )

    started_at = time.perf_counter()

    answer, tool_calls, sources, trace = (
        run_knowledge_base_agent(
            session=session,
            knowledge_base_id=(
                knowledge_base_id
            ),
            repository_analysis_task_id=(
                repository_scope_id
            ),
            question=question,
            top_k=request.top_k,
            max_steps=request.max_steps,
            semantic_weight=(
                request.semantic_weight
            ),
            keyword_weight=(
                request.keyword_weight
            ),
        )
    )

    if repository_scope_id is not None:
        invalid_sources = [
            source
            for source in sources
            if (
                    source.repository_analysis_task_id
                    != repository_scope_id
            )
        ]

        if invalid_sources:
            invalid_source_details = [
                {
                    "chunk_id": str(
                        source.chunk_id,
                    ),
                    "document_id": str(
                        source.document_id,
                    ),
                    "repository_analysis_task_id": (
                        str(
                            source.repository_analysis_task_id,
                        )
                        if (
                                source.repository_analysis_task_id
                                is not None
                        )
                        else None
                    ),
                    "original_filename": (
                        source.original_filename
                    ),
                    "repository_relative_path": (
                        source.repository_relative_path
                    ),
                }
                for source in invalid_sources[
                              :5
                              ]
            ]

            raise HTTPException(
                status_code=500,
                detail={
                    "code": (
                        "REPOSITORY_AGENT_"
                        "SCOPE_VIOLATION"
                    ),
                    "message": (
                        "Repository-scoped Agent "
                        "returned sources outside "
                        "the requested repository "
                        "analysis task"
                    ),
                    "invalid_sources": (
                        invalid_source_details
                    ),
                },
            )

        trace.append(
            "repository_scope_sources_"
            f"validated={len(sources)}",
        )

    latency_ms = int(
        (
                time.perf_counter()
                - started_at
        )
        * 1000
    )

    agent_run = AgentRun(
        knowledge_base_id=(
            knowledge_base_id
        ),

        owner_id=current_user.id,

        repository_analysis_task_id=(
            repository_scope_id
        ),

        source_commit_sha=(
            repository_scope_task
                .resolved_commit_sha
            if repository_scope_task
               is not None
            else None
        ),

        question=question,
        answer=answer,

        top_k=request.top_k,
        max_steps=request.max_steps,

        semantic_weight=(
            request.semantic_weight
        ),

        keyword_weight=(
            request.keyword_weight
        ),

        latency_ms=latency_ms,

        tool_calls_json=(
            agent_tool_calls_to_json(
                tool_calls,
            )
        ),

        sources_json=(
            agent_sources_to_json(
                sources,
            )
        ),

        trace_json=(
            agent_trace_to_json(
                trace,
            )
        ),
    )

    session.add(agent_run)
    session.commit()
    session.refresh(agent_run)

    return AgentChatResponse(
        agent_run_id=agent_run.id,
        answer=answer,
        tool_calls=tool_calls,
        sources=sources,
        trace=trace,
        latency_ms=latency_ms,
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/agent-runs",
    response_model=AgentRunsPublic,
    operation_id="read_agent_runs",
)
def read_agent_runs(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 20,
    keyword: str | None = None,
    repository_analysis_task_id: uuid.UUID | None = None,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )
    repository_scope_id = (
        resolve_repository_analysis_scope(
            session=session,
            current_user=current_user,
            knowledge_base_id=(
                knowledge_base_id
            ),
            repository_analysis_task_id=(
                repository_analysis_task_id
            ),
        )
    )

    filters = [
        AgentRun.knowledge_base_id
        == knowledge_base_id,
    ]

    if repository_scope_id is None:
        filters.append(
            col(
                AgentRun
                    .repository_analysis_task_id
            ).is_(
                None,
            ),
        )
    else:
        filters.append(
            AgentRun
            .repository_analysis_task_id
            == repository_scope_id,
        )

    if keyword:
        keyword_pattern = f"%{keyword.strip()}%"
        filters.append(
            col(AgentRun.question).ilike(keyword_pattern)
            | col(AgentRun.answer).ilike(keyword_pattern)
        )

    count_statement = (
        select(func.count())
        .select_from(AgentRun)
        .where(*filters)
    )

    count = session.exec(count_statement).one()

    statement = (
        select(AgentRun)
        .where(*filters)
        .order_by(col(AgentRun.created_at).desc())
        .offset(skip)
        .limit(limit)
    )

    agent_runs = session.exec(statement).all()

    return AgentRunsPublic(
        data=[agent_run_to_public(agent_run) for agent_run in agent_runs],
        count=count,
    )



@router.delete(
    "/agent-runs/{agent_run_id}",
    response_model=Message,
    operation_id="delete_agent_run",
)
def delete_agent_run(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    agent_run_id: uuid.UUID,
) -> Any:
    agent_run = session.get(AgentRun, agent_run_id)

    if not agent_run:
        raise HTTPException(status_code=404, detail="Agent run not found")

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=agent_run.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    session.delete(agent_run)
    session.commit()

    return Message(message="Agent run deleted successfully")


@router.get(
    "/knowledge-bases/{knowledge_base_id}/agent-settings",
    response_model=KnowledgeBaseAgentSettingsPublic,
    operation_id="read_knowledge_base_agent_settings",
)
def read_knowledge_base_agent_settings(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    settings = get_or_create_agent_settings(
        session=session,
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
    )

    return build_agent_settings_public(settings)


@router.put(
    "/knowledge-bases/{knowledge_base_id}/agent-settings",
    response_model=KnowledgeBaseAgentSettingsPublic,
    operation_id="update_knowledge_base_agent_settings",
)
def update_knowledge_base_agent_settings(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: KnowledgeBaseAgentSettingsUpdate,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    settings = get_or_create_agent_settings(
        session=session,
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
    )

    next_top_k = request.top_k if request.top_k is not None else settings.top_k
    next_max_steps = (
        request.max_steps
        if request.max_steps is not None
        else settings.max_steps
    )
    next_semantic_weight = (
        request.semantic_weight
        if request.semantic_weight is not None
        else settings.semantic_weight
    )
    next_keyword_weight = (
        request.keyword_weight
        if request.keyword_weight is not None
        else settings.keyword_weight
    )

    next_semantic_weight, next_keyword_weight = normalize_retrieval_weights(
        semantic_weight=next_semantic_weight,
        keyword_weight=next_keyword_weight,
    )

    settings.top_k = next_top_k
    settings.max_steps = next_max_steps
    settings.semantic_weight = next_semantic_weight
    settings.keyword_weight = next_keyword_weight
    settings.updated_at = get_datetime_utc()

    session.add(settings)
    session.commit()
    session.refresh(settings)

    return build_agent_settings_public(settings)


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-runs",
    response_model=RagRunsPublic,
    operation_id="read_knowledge_base_rag_runs",
)
def read_knowledge_base_rag_runs(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 20,
    keyword: str | None = None,
    repository_analysis_task_id: uuid.UUID | None = None,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )
    repository_scope_id = (
        resolve_repository_analysis_scope(
            session=session,
            current_user=current_user,
            knowledge_base_id=(
                knowledge_base_id
            ),
            repository_analysis_task_id=(
                repository_analysis_task_id
            ),
        )
    )
    filters = [
        RagRun.knowledge_base_id
        == knowledge_base_id,
    ]

    if repository_scope_id is None:
        filters.append(
            col(
                RagRun
                    .repository_analysis_task_id
            ).is_(
                None,
            ),
        )
    else:
        filters.append(
            RagRun
            .repository_analysis_task_id
            == repository_scope_id,
        )

    if keyword and keyword.strip():
        pattern = f"%{keyword.strip()}%"
        filters.append(
            or_(
                col(RagRun.question).ilike(pattern),
                col(RagRun.answer).ilike(pattern),
                col(RagRun.retrieval_type).ilike(pattern),
            )
        )

    count_statement = (
        select(func.count())
        .select_from(RagRun)
        .where(*filters)
    )
    count = session.exec(count_statement).one()

    statement = (
        select(RagRun)
        .where(*filters)
        .order_by(col(RagRun.created_at).desc())
        .offset(skip)
        .limit(limit)
    )

    rag_runs = session.exec(statement).all()

    return RagRunsPublic(
        data=[rag_run_to_public(rag_run) for rag_run in rag_runs],
        count=count,
    )


@router.delete(
    "/rag-runs/{rag_run_id}",
    response_model=Message,
    operation_id="delete_rag_run",
)
def delete_rag_run(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    rag_run_id: uuid.UUID,
) -> Any:
    rag_run = session.get(RagRun, rag_run_id)

    if not rag_run:
        raise HTTPException(status_code=404, detail="RAG run not found")

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=rag_run.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    session.delete(rag_run)
    session.commit()

    return Message(message="RAG run deleted successfully")


@router.post(
    "/knowledge-bases/{knowledge_base_id}/embeddings/backfill",
    response_model=KnowledgeBaseEmbeddingBackfillResult,
    operation_id="backfill_knowledge_base_embeddings",
)
def backfill_knowledge_base_embeddings(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: KnowledgeBaseEmbeddingBackfillRequest,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    result = backfill_chunk_embeddings(
        session=session,
        knowledge_base_id=knowledge_base_id,
        limit=request.limit,
        retry_failed=request.retry_failed,
        force=request.force,
    )

    return KnowledgeBaseEmbeddingBackfillResult(**result)


@router.post(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-cases",
    response_model=RagEvalCasePublic,
    operation_id="create_rag_eval_case",
)
def create_rag_eval_case(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: RagEvalCaseCreate,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    question = request.question.strip()

    if not question:
        raise HTTPException(status_code=400, detail="Question is required")

    eval_case = RagEvalCase(
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
        question=question,
        expected_keywords_json=json.dumps(
            request.expected_keywords,
            ensure_ascii=False,
        ),
        expected_source_filename=request.expected_source_filename,
        note=request.note,
    )

    session.add(eval_case)
    session.commit()
    session.refresh(eval_case)

    return rag_eval_case_to_public(eval_case)


@router.post(
    "/knowledge-bases/{knowledge_base_id}/code-eval-cases/seed",
    response_model=RagEvalCasesPublic,
    operation_id="seed_code_eval_cases",
)
def seed_code_eval_cases(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    max_files: int = 20,
    max_symbols: int = 40,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )

    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    max_files = max(1, min(max_files, 100))
    max_symbols = max(1, min(max_symbols, 200))

    created_cases = build_code_eval_cases_from_documents(
        session=session,
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
        max_files=max_files,
        max_symbols=max_symbols,
    )

    if not created_cases:
        raise HTTPException(
            status_code=400,
            detail="当前知识库中没有可用于生成 Code Eval Cases 的代码文件或代码 chunk。",
        )

    session.commit()

    for eval_case in created_cases:
        session.refresh(eval_case)

    return RagEvalCasesPublic(
        data=[
            rag_eval_case_to_public(eval_case)
            for eval_case in created_cases
        ],
        count=len(created_cases),
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/code-eval/type-compare-summary",
    response_model=CodeEvalTypeCompareSummaryPublic,
    operation_id="get_code_eval_type_compare_summary",
)
def get_code_eval_type_compare_summary(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )

    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    cases_statement = (
        select(RagEvalCase)
        .where(RagEvalCase.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalCase.created_at).desc())
    )

    eval_cases = session.exec(cases_statement).all()

    stats_by_type: dict[str, dict[str, Any]] = {}

    for eval_case in eval_cases:
        case_type = parse_code_eval_case_type(eval_case.note)

        if not case_type.startswith("file_") and case_type not in {
            "symbol_definition",
            "symbol_reference",
            "code_logic",
            "bug_location",
        }:
            continue

        if case_type not in stats_by_type:
            stats_by_type[case_type] = {
                "case_type": case_type,
                "case_type_label": get_code_eval_case_type_label(case_type),
                "total_cases": 0,
                "compared_cases": 0,
                "rag_source_hits": 0,
                "agent_source_hits": 0,
                "rag_keyword_hit_sum": 0.0,
                "agent_keyword_hit_sum": 0.0,
                "rag_latency_sum": 0.0,
                "rag_latency_count": 0,
                "agent_latency_sum": 0.0,
                "agent_latency_count": 0,
                "agent_tool_call_sum": 0,
                "agent_tool_call_count": 0,
                "agent_tool_usage": {},
            }

        stats = stats_by_type[case_type]
        stats["total_cases"] += 1

        compare_statement = (
            select(RagAgentCompareItem)
                .where(RagAgentCompareItem.eval_case_id == eval_case.id)
                .order_by(col(RagAgentCompareItem.created_at).desc())
                .limit(1)
        )

        compare_run = session.exec(compare_statement).first()

        if not compare_run:
            continue

        stats["compared_cases"] += 1

        rag_answer = getattr(compare_run, "rag_answer", None)
        agent_answer = getattr(compare_run, "agent_answer", None)

        rag_sources_json = getattr(compare_run, "rag_sources_json", None)
        agent_sources_json = getattr(compare_run, "agent_sources_json", None)

        rag_latency_ms = getattr(compare_run, "rag_latency_ms", None)
        agent_latency_ms = getattr(compare_run, "agent_latency_ms", None)

        agent_tool_calls_json = getattr(compare_run, "agent_tool_calls_json", None)

        if calculate_source_file_hit(
            sources_value=rag_sources_json,
            expected_source_filename=eval_case.expected_source_filename,
        ):
            stats["rag_source_hits"] += 1

        if calculate_source_file_hit(
            sources_value=agent_sources_json,
            expected_source_filename=eval_case.expected_source_filename,
        ):
            stats["agent_source_hits"] += 1

        stats["rag_keyword_hit_sum"] += calculate_keyword_hit_rate(
            answer=rag_answer,
            expected_keywords_json=eval_case.expected_keywords_json,
        )

        stats["agent_keyword_hit_sum"] += calculate_keyword_hit_rate(
            answer=agent_answer,
            expected_keywords_json=eval_case.expected_keywords_json,
        )

        if rag_latency_ms is not None:
            stats["rag_latency_sum"] += float(rag_latency_ms)
            stats["rag_latency_count"] += 1

        if agent_latency_ms is not None:
            stats["agent_latency_sum"] += float(agent_latency_ms)
            stats["agent_latency_count"] += 1

        tool_call_count, tool_usage = count_agent_tool_calls(agent_tool_calls_json)

        stats["agent_tool_call_sum"] += tool_call_count
        stats["agent_tool_call_count"] += 1

        merge_tool_usage(
            target=stats["agent_tool_usage"],
            current=tool_usage,
        )

    result: list[CodeEvalTypeCompareStat] = []

    order = [
        "file_overview",
        "symbol_definition",
        "symbol_reference",
        "code_logic",
        "bug_location",
        "unknown",
    ]

    for case_type in order:
        if case_type not in stats_by_type:
            continue

        item = stats_by_type[case_type]
        compared_cases = item["compared_cases"]

        if compared_cases > 0:
            rag_source_hit_rate = item["rag_source_hits"] / compared_cases
            agent_source_hit_rate = item["agent_source_hits"] / compared_cases
            rag_keyword_hit_rate = item["rag_keyword_hit_sum"] / compared_cases
            agent_keyword_hit_rate = item["agent_keyword_hit_sum"] / compared_cases
        else:
            rag_source_hit_rate = 0
            agent_source_hit_rate = 0
            rag_keyword_hit_rate = 0
            agent_keyword_hit_rate = 0

        rag_avg_latency_ms = (
            item["rag_latency_sum"] / item["rag_latency_count"]
            if item["rag_latency_count"] > 0
            else None
        )

        agent_avg_latency_ms = (
            item["agent_latency_sum"] / item["agent_latency_count"]
            if item["agent_latency_count"] > 0
            else None
        )

        agent_avg_tool_calls = (
            item["agent_tool_call_sum"] / item["agent_tool_call_count"]
            if item["agent_tool_call_count"] > 0
            else None
        )

        result.append(
            CodeEvalTypeCompareStat(
                case_type=item["case_type"],
                case_type_label=item["case_type_label"],
                total_cases=item["total_cases"],
                compared_cases=compared_cases,
                rag_source_hit_rate=rag_source_hit_rate,
                agent_source_hit_rate=agent_source_hit_rate,
                rag_keyword_hit_rate=rag_keyword_hit_rate,
                agent_keyword_hit_rate=agent_keyword_hit_rate,
                rag_avg_latency_ms=rag_avg_latency_ms,
                agent_avg_latency_ms=agent_avg_latency_ms,
                agent_avg_tool_calls=agent_avg_tool_calls,
                agent_tool_usage_json=json.dumps(
                    item["agent_tool_usage"],
                    ensure_ascii=False,
                ),
            )
        )

    return CodeEvalTypeCompareSummaryPublic(
        data=result,
        count=len(result),
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-agent-compare-failure-analysis",
    response_model=RagAgentCompareFailureAnalysisPublic,
    operation_id="read_rag_agent_compare_failure_analysis",
)
def read_rag_agent_compare_failure_analysis(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    batch_id: uuid.UUID | None = None,
    case_type: str | None = None,
    limit: int = 50,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )

    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    limit = max(1, min(limit, 200))

    filters = [
        RagAgentCompareItem.knowledge_base_id == knowledge_base_id,
    ]

    if batch_id is not None:
        filters.append(RagAgentCompareItem.batch_id == batch_id)

    statement = (
        select(RagAgentCompareItem)
        .where(*filters)
        .order_by(col(RagAgentCompareItem.created_at).desc())
        .limit(limit * 3)
    )

    compare_items = session.exec(statement).all()

    eval_case_ids = [
        item.eval_case_id
        for item in compare_items
        if item.eval_case_id is not None
    ]

    eval_cases_by_id: dict[uuid.UUID, RagEvalCase] = {}

    if eval_case_ids:
        eval_cases_statement = (
            select(RagEvalCase)
            .where(col(RagEvalCase.id).in_(eval_case_ids))
        )

        eval_cases = session.exec(eval_cases_statement).all()

        eval_cases_by_id = {
            eval_case.id: eval_case
            for eval_case in eval_cases
        }

    failure_cases: list[RagAgentCompareFailureCasePublic] = []
    reason_count_map: dict[str, int] = {}

    total_items = len(compare_items)

    for item in compare_items:
        eval_case = None

        if item.eval_case_id is not None:
            eval_case = eval_cases_by_id.get(item.eval_case_id)

        current_case_type = "unknown"

        if eval_case is not None:
            current_case_type = parse_code_eval_case_type(eval_case.note)

        if case_type and current_case_type != case_type:
            continue

        failure_reasons = build_compare_failure_reasons(item=item)

        if not failure_reasons:
            continue

        failure_case = build_compare_failure_case_public(
            item=item,
            eval_case=eval_case,
        )

        failure_cases.append(failure_case)

        for reason in failure_reasons:
            reason_count_map[reason] = reason_count_map.get(reason, 0) + 1

        if len(failure_cases) >= limit:
            break

    reason_stats = [
        RagAgentCompareFailureReasonStat(
            reason=reason,
            reason_label=get_compare_failure_reason_label(reason),
            count=count,
        )
        for reason, count in sorted(
            reason_count_map.items(),
            key=lambda item: item[1],
            reverse=True,
        )
    ]

    return RagAgentCompareFailureAnalysisPublic(
        total_items=total_items,
        failed_items=len(failure_cases),
        reason_stats=reason_stats,
        data=failure_cases,
        count=len(failure_cases),
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-cases",
    response_model=RagEvalCasesPublic,
    operation_id="read_rag_eval_cases",
)
def read_rag_eval_cases(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 20,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    count_statement = (
        select(func.count())
        .select_from(RagEvalCase)
        .where(RagEvalCase.knowledge_base_id == knowledge_base_id)
    )
    count = session.exec(count_statement).one()

    statement = (
        select(RagEvalCase)
        .where(RagEvalCase.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalCase.created_at).desc())
        .offset(skip)
        .limit(limit)
    )

    eval_cases = session.exec(statement).all()

    return RagEvalCasesPublic(
        data=[rag_eval_case_to_public(eval_case) for eval_case in eval_cases],
        count=count,
    )


@router.delete(
    "/rag-eval-cases/{eval_case_id}",
    response_model=Message,
    operation_id="delete_rag_eval_case",
)
def delete_rag_eval_case(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    eval_case_id: uuid.UUID,
) -> Any:
    eval_case = session.get(RagEvalCase, eval_case_id)

    if not eval_case:
        raise HTTPException(status_code=404, detail="RAG eval case not found")

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=eval_case.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    session.delete(eval_case)
    session.commit()

    return Message(message="RAG eval case deleted successfully")


def run_eval_case_and_save(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    eval_case: RagEvalCase,
    top_k: int,
    semantic_weight: float,
    keyword_weight: float,
    batch_id: uuid.UUID | None = None,
) -> RagEvalRun:
    started_at = time.perf_counter()
    trace: list[str] = ["eval_case_loaded"]

    trace.append(
        f"retrieval_params_top_k={top_k},semantic_weight={semantic_weight:.2f},keyword_weight={keyword_weight:.2f}"
    )

    hybrid_rows = hybrid_retrieve_chunks_for_rag(
        session=session,
        knowledge_base_id=eval_case.knowledge_base_id,
        query=eval_case.question,
        top_k=top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
    )

    trace.append("hybrid_retrieve_chunks")

    sources = [
        RagChatSource(
            document_id=document.id,
            original_filename=document.original_filename,
            chunk_id=chunk.id,
            chunk_index=chunk.chunk_index,
            content=chunk.content,
            content_length=chunk.content_length,
            match_count=match_count,
            similarity=similarity,
            retrieval_type="hybrid",
        )
        for chunk, document, similarity, match_count, hybrid_score in hybrid_rows
    ]

    trace.append("build_hybrid_sources")

    answer = generate_rag_answer(
        question=eval_case.question,
        sources=sources,
    )

    trace.append("generate_answer")

    latency_ms = int((time.perf_counter() - started_at) * 1000)

    expected_keywords = parse_json_list(eval_case.expected_keywords_json)

    keyword_hit = check_keyword_hit(
        answer=answer,
        expected_keywords=expected_keywords,
    )

    source_hit = check_source_hit(
        sources=sources,
        expected_source_filename=eval_case.expected_source_filename,
    )

    score = calculate_eval_score(
        keyword_hit=keyword_hit,
        source_hit=source_hit,
    )

    trace.append("calculate_eval_score")

    eval_run = RagEvalRun(
        eval_case_id=eval_case.id,
        batch_id=batch_id,
        knowledge_base_id=eval_case.knowledge_base_id,
        owner_id=current_user.id,
        question=eval_case.question,
        answer=answer,
        expected_keywords_json=eval_case.expected_keywords_json,
        expected_source_filename=eval_case.expected_source_filename,
        sources_json=json.dumps(
            jsonable_encoder(sources),
            ensure_ascii=False,
        ),
        trace_json=json.dumps(
            trace,
            ensure_ascii=False,
        ),
        keyword_hit=keyword_hit,
        source_hit=source_hit,
        score=score,
        retrieval_type="hybrid",
        top_k=top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
        latency_ms=latency_ms,
    )

    session.add(eval_run)
    session.commit()
    session.refresh(eval_run)

    return eval_run


@router.post(
    "/rag-eval-cases/{eval_case_id}/run",
    response_model=RagEvalRunPublic,
    operation_id="run_rag_eval_case",
)
def run_rag_eval_case(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    eval_case_id: uuid.UUID,
    request: RagEvalRunRequest,
) -> Any:
    eval_case = session.get(RagEvalCase, eval_case_id)

    if not eval_case:
        raise HTTPException(status_code=404, detail="RAG eval case not found")

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=eval_case.knowledge_base_id,
    )

    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    semantic_weight, keyword_weight = normalize_retrieval_weights(
        semantic_weight=request.semantic_weight,
        keyword_weight=request.keyword_weight,
    )

    eval_run = run_eval_case_and_save(
        session=session,
        current_user=current_user,
        eval_case=eval_case,
        top_k=request.top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
    )

    return rag_eval_run_to_public(eval_run)


@router.post(
    "/knowledge-bases/{knowledge_base_id}/rag-retrieval-presets",
    response_model=RagRetrievalPresetPublic,
    operation_id="create_rag_retrieval_preset",
)
def create_rag_retrieval_preset(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: RagRetrievalPresetCreate,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    semantic_weight, keyword_weight = normalize_retrieval_weights(
        semantic_weight=request.semantic_weight,
        keyword_weight=request.keyword_weight,
    )

    preset = RagRetrievalPreset(
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
        name=request.name.strip(),
        top_k=request.top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
        note=request.note,
    )

    session.add(preset)
    session.commit()
    session.refresh(preset)

    return rag_retrieval_preset_to_public(preset)


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-retrieval-presets",
    response_model=RagRetrievalPresetsPublic,
    operation_id="read_rag_retrieval_presets",
)
def read_rag_retrieval_presets(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 50,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    count_statement = (
        select(func.count())
        .select_from(RagRetrievalPreset)
        .where(RagRetrievalPreset.knowledge_base_id == knowledge_base_id)
    )
    count = session.exec(count_statement).one()

    statement = (
        select(RagRetrievalPreset)
        .where(RagRetrievalPreset.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagRetrievalPreset.created_at).desc())
        .offset(skip)
        .limit(limit)
    )

    presets = session.exec(statement).all()

    return RagRetrievalPresetsPublic(
        data=[rag_retrieval_preset_to_public(preset) for preset in presets],
        count=count,
    )


@router.delete(
    "/rag-retrieval-presets/{preset_id}",
    response_model=Message,
    operation_id="delete_rag_retrieval_preset",
)
def delete_rag_retrieval_preset(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    preset_id: uuid.UUID,
) -> Any:
    preset = session.get(RagRetrievalPreset, preset_id)

    if not preset:
        raise HTTPException(status_code=404, detail="RAG retrieval preset not found")

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=preset.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    session.delete(preset)
    session.commit()

    return Message(message="RAG retrieval preset deleted successfully")


@router.post(
    "/knowledge-bases/{knowledge_base_id}/rag-agent-compare-eval",
    response_model=RagAgentCompareEvalResponse,
    operation_id="run_rag_agent_compare_eval",
)
def run_rag_agent_compare_eval(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: RagAgentCompareEvalRequest,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    semantic_weight, keyword_weight = normalize_retrieval_weights(
        semantic_weight=request.semantic_weight,
        keyword_weight=request.keyword_weight,
    )

    statement = (
        select(RagEvalCase)
        .where(RagEvalCase.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalCase.created_at).asc())
        .limit(request.limit)
    )

    eval_cases = session.exec(statement).all()

    items: list[RagAgentCompareEvalItem] = []

    for eval_case in eval_cases:
        expected_keywords = parse_json_list(eval_case.expected_keywords_json)

        rag_answer: str | None = None
        rag_sources: list[RagChatSource] = []
        rag_trace: list[str] = []
        rag_latency_ms: int | None = None
        rag_error_message: str | None = None

        agent_answer: str | None = None
        agent_tool_calls: list[AgentToolCallPublic] = []
        agent_sources: list[RagChatSource] = []
        agent_trace: list[str] = []
        agent_latency_ms: int | None = None
        agent_error_message: str | None = None

        try:
            rag_answer, rag_sources, rag_trace, rag_latency_ms = (
                run_rag_once_for_compare(
                    session=session,
                    knowledge_base_id=knowledge_base_id,
                    question=eval_case.question,
                    top_k=request.top_k,
                    semantic_weight=semantic_weight,
                    keyword_weight=keyword_weight,
                )
            )
        except Exception as error:
            rag_error_message = f"{error.__class__.__name__}: {error}"
            rag_trace.append("compare_rag_failed")

        try:
            (
                agent_answer,
                agent_tool_calls,
                agent_sources,
                agent_trace,
                agent_latency_ms,
            ) = run_agent_once_for_compare(
                session=session,
                knowledge_base_id=knowledge_base_id,
                question=eval_case.question,
                top_k=request.top_k,
                max_steps=request.max_steps,
                semantic_weight=semantic_weight,
                keyword_weight=keyword_weight,
            )
        except Exception as error:
            agent_error_message = f"{error.__class__.__name__}: {error}"
            agent_trace.append("compare_agent_failed")

        agent_failed_tool_count = sum(
            1 for tool_call in agent_tool_calls if not tool_call.success
        )

        agent_tool_names = [
            tool_call.tool_name for tool_call in agent_tool_calls
        ]

        is_failed = (
            rag_error_message is not None
            or agent_error_message is not None
            or agent_failed_tool_count > 0
        )

        items.append(
            RagAgentCompareEvalItem(
                eval_case_id=eval_case.id,
                question=eval_case.question,
                expected_keywords=expected_keywords,
                expected_source_filename=eval_case.expected_source_filename,
                rag_answer=rag_answer,
                agent_answer=agent_answer,
                rag_latency_ms=rag_latency_ms,
                agent_latency_ms=agent_latency_ms,
                rag_sources_count=len(rag_sources),
                agent_sources_count=len(agent_sources),
                agent_tool_call_count=len(agent_tool_calls),
                agent_failed_tool_count=agent_failed_tool_count,
                agent_tool_names=agent_tool_names,
                rag_sources=rag_sources,
                agent_sources=agent_sources,
                agent_tool_calls=agent_tool_calls,
                rag_trace=rag_trace,
                agent_trace=agent_trace,
                rag_error_message=rag_error_message,
                agent_error_message=agent_error_message,
                is_failed=is_failed,
            )
        )

    summary = build_rag_agent_compare_summary(items)

    batch = save_rag_agent_compare_batch(
        session=session,
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
        request=request,
        summary=summary,
        items=items,
    )

    return RagAgentCompareEvalResponse(
        compare_batch_id=batch.id,
        summary=summary,
        data=items,
        count=len(items),
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-agent-compare-batches",
    response_model=RagAgentCompareBatchesPublic,
    operation_id="read_rag_agent_compare_batches",
)
def read_rag_agent_compare_batches(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 20,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    count_statement = (
        select(func.count())
        .select_from(RagAgentCompareBatch)
        .where(RagAgentCompareBatch.knowledge_base_id == knowledge_base_id)
    )

    count = session.exec(count_statement).one()

    statement = (
        select(RagAgentCompareBatch)
        .where(RagAgentCompareBatch.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagAgentCompareBatch.created_at).desc())
        .offset(skip)
        .limit(limit)
    )

    batches = session.exec(statement).all()

    return RagAgentCompareBatchesPublic(
        data=[
            build_compare_batch_public(batch)
            for batch in batches
        ],
        count=count,
    )


@router.get(
    "/rag-agent-compare-batches/{batch_id}",
    response_model=RagAgentCompareBatchDetailPublic,
    operation_id="read_rag_agent_compare_batch",
)
def read_rag_agent_compare_batch(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    batch_id: uuid.UUID,
) -> Any:
    batch = session.get(RagAgentCompareBatch, batch_id)

    if not batch:
        raise HTTPException(
            status_code=404,
            detail="RAG vs Agent compare batch not found",
        )

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=batch.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = (
        select(RagAgentCompareItem)
        .where(RagAgentCompareItem.batch_id == batch_id)
        .order_by(col(RagAgentCompareItem.created_at).asc())
    )

    items = session.exec(statement).all()

    return RagAgentCompareBatchDetailPublic(
      batch=build_compare_batch_public(batch),
      data=[
        build_compare_item_public(item)
        for item in items
      ],
      count=len(items),
    )



@router.get(
    "/knowledge-bases/{knowledge_base_id}/code-agent-experiment-report",
    response_model=CodeAgentExperimentReportPublic,
    operation_id="read_code_agent_experiment_report",
)
def read_code_agent_experiment_report(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    baseline_batch_id: uuid.UUID,
    optimized_batch_id: uuid.UUID,
    failure_limit: int = 8,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )

    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    baseline_batch = session.get(RagAgentCompareBatch, baseline_batch_id)

    if (
        baseline_batch is None
        or baseline_batch.knowledge_base_id != knowledge_base_id
    ):
        raise HTTPException(
            status_code=404,
            detail="Baseline compare batch not found",
        )

    optimized_batch = session.get(RagAgentCompareBatch, optimized_batch_id)

    if (
        optimized_batch is None
        or optimized_batch.knowledge_base_id != knowledge_base_id
    ):
        raise HTTPException(
            status_code=404,
            detail="Optimized compare batch not found",
        )

    failure_limit = max(1, min(failure_limit, 30))

    metric_rows = build_compare_batch_metric_rows(
        baseline_batch=baseline_batch,
        optimized_batch=optimized_batch,
    )

    type_stats = build_code_eval_type_stats_for_compare_batch(
        session=session,
        batch_id=optimized_batch.id,
    )

    reason_count_map, failure_items = build_compare_failure_report_data(
        session=session,
        batch_id=optimized_batch.id,
        limit=failure_limit,
    )

    advice = build_backend_code_agent_advice(
        reason_count_map=reason_count_map,
        type_stats=type_stats,
    )

    content = build_code_agent_experiment_report_markdown(
        knowledge_base=knowledge_base,
        baseline_batch=baseline_batch,
        optimized_batch=optimized_batch,
        metric_rows=metric_rows,
        type_stats=type_stats,
        reason_count_map=reason_count_map,
        failure_items=failure_items,
        advice=advice,
    )

    safe_name = "".join(
        char if char.isalnum() or char in {"-", "_"} else "_"
        for char in knowledge_base.name
    ).strip("_") or "knowledge_base"

    return CodeAgentExperimentReportPublic(
        filename=f"code_agent_experiment_report_{safe_name}.md",
        mime_type="text/markdown",
        content=content,
    )


@router.get(
    "/rag-agent-compare-batches/{batch_id}/report",
    response_model=RagAgentCompareReportResponse,
    operation_id="read_rag_agent_compare_batch_report",
)
def read_rag_agent_compare_batch_report(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    batch_id: uuid.UUID,
) -> Any:
    batch = session.get(RagAgentCompareBatch, batch_id)

    if not batch:
        raise HTTPException(
            status_code=404,
            detail="RAG vs Agent compare batch not found",
        )

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=batch.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = (
        select(RagAgentCompareItem)
        .where(RagAgentCompareItem.batch_id == batch_id)
        .order_by(col(RagAgentCompareItem.created_at).asc())
    )

    db_items = session.exec(statement).all()

    batch_public = build_compare_batch_public(batch)
    item_publics = [
        build_compare_item_public(item)
        for item in db_items
    ]

    knowledge_base_name = getattr(
        knowledge_base,
        "name",
        str(knowledge_base.id),
    )

    content = build_rag_agent_compare_markdown_report(
        knowledge_base_name=knowledge_base_name,
        batch=batch_public,
        items=item_publics,
    )

    safe_batch_name = "".join(
        char if char.isalnum() or char in ["-", "_"] else "_"
        for char in batch.name
    )

    filename = f"rag_agent_compare_report_{safe_batch_name}.md"

    return RagAgentCompareReportResponse(
        filename=filename,
        content=content,
    )


@router.get(
    "/rag-agent-compare-batches/{batch_id}/report-docx",
    response_model=RagAgentCompareFileReportResponse,
    operation_id="read_rag_agent_compare_batch_docx_report",
)
def read_rag_agent_compare_batch_docx_report(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    batch_id: uuid.UUID,
) -> Any:
    batch = session.get(RagAgentCompareBatch, batch_id)

    if not batch:
        raise HTTPException(
            status_code=404,
            detail="RAG vs Agent compare batch not found",
        )

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=batch.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = (
        select(RagAgentCompareItem)
        .where(RagAgentCompareItem.batch_id == batch_id)
        .order_by(col(RagAgentCompareItem.created_at).asc())
    )

    db_items = session.exec(statement).all()

    batch_public = build_compare_batch_public(batch)
    item_publics = [
        build_compare_item_public(item)
        for item in db_items
    ]

    knowledge_base_name = getattr(
        knowledge_base,
        "name",
        str(knowledge_base.id),
    )

    file_bytes = build_rag_agent_compare_docx_report(
        knowledge_base_name=knowledge_base_name,
        batch=batch_public,
        items=item_publics,
    )

    safe_batch_name = "".join(
        char if char.isalnum() or char in ["-", "_"] else "_"
        for char in batch.name
    )

    filename = f"rag_agent_compare_report_{safe_batch_name}.docx"

    return build_base64_file_response(
        filename=filename,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_bytes=file_bytes,
    )

@router.get(
    "/rag-agent-compare-batches/{batch_id}/report-pdf",
    response_model=RagAgentCompareFileReportResponse,
    operation_id="read_rag_agent_compare_batch_pdf_report",
)
def read_rag_agent_compare_batch_pdf_report(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    batch_id: uuid.UUID,
) -> Any:
    batch = session.get(RagAgentCompareBatch, batch_id)

    if not batch:
        raise HTTPException(
            status_code=404,
            detail="RAG vs Agent compare batch not found",
        )

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=batch.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = (
        select(RagAgentCompareItem)
        .where(RagAgentCompareItem.batch_id == batch_id)
        .order_by(col(RagAgentCompareItem.created_at).asc())
    )

    db_items = session.exec(statement).all()

    batch_public = build_compare_batch_public(batch)
    item_publics = [
        build_compare_item_public(item)
        for item in db_items
    ]

    knowledge_base_name = getattr(
        knowledge_base,
        "name",
        str(knowledge_base.id),
    )

    file_bytes = build_rag_agent_compare_pdf_report(
        knowledge_base_name=knowledge_base_name,
        batch=batch_public,
        items=item_publics,
    )

    safe_batch_name = "".join(
        char if char.isalnum() or char in ["-", "_"] else "_"
        for char in batch.name
    )

    filename = f"rag_agent_compare_report_{safe_batch_name}.pdf"

    return build_base64_file_response(
        filename=filename,
        mime_type="application/pdf",
        file_bytes=file_bytes,
    )

@router.delete(
    "/rag-agent-compare-batches/{batch_id}",
    operation_id="delete_rag_agent_compare_batch",
)
def delete_rag_agent_compare_batch(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    batch_id: uuid.UUID,
) -> Message:
    batch = session.get(RagAgentCompareBatch, batch_id)

    if not batch:
        raise HTTPException(
            status_code=404,
            detail="RAG vs Agent compare batch not found",
        )

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=batch.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    session.delete(batch)
    session.commit()

    return Message(message="RAG vs Agent compare batch deleted")


@router.post(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-cases/run-all",
    response_model=RagEvalBatchRunResult,
    operation_id="run_all_rag_eval_cases",
)
def run_all_rag_eval_cases(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    request: RagEvalBatchRunRequest,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    semantic_weight, keyword_weight = normalize_retrieval_weights(
        semantic_weight=request.semantic_weight,
        keyword_weight=request.keyword_weight,
    )

    statement = (
        select(RagEvalCase)
        .where(RagEvalCase.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalCase.created_at).asc())
        .limit(request.limit)
    )

    eval_cases = session.exec(statement).all()

    batch_name = request.batch_name.strip() if request.batch_name else ""

    if not batch_name:
        batch_name = (
            f"Batch top_k={request.top_k} "
            f"S={semantic_weight:.2f} K={keyword_weight:.2f}"
        )

    batch = RagEvalBatch(
        knowledge_base_id=knowledge_base_id,
        owner_id=current_user.id,
        name=batch_name,
        note=request.batch_note,
        top_k=request.top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
        total_cases=len(eval_cases),
        ran=0,
        failed=0,
    )

    session.add(batch)
    session.commit()
    session.refresh(batch)

    eval_runs: list[RagEvalRun] = []
    failed = 0

    for eval_case in eval_cases:
        try:
            eval_run = run_eval_case_and_save(
                session=session,
                current_user=current_user,
                eval_case=eval_case,
                top_k=request.top_k,
                semantic_weight=semantic_weight,
                keyword_weight=keyword_weight,
                batch_id=batch.id,
            )
            eval_runs.append(eval_run)
        except Exception as error:
            failed += 1
            print(f"Run eval case failed: {eval_case.id}, error: {error}")

    summary = build_eval_summary_from_runs(eval_runs)

    batch.ran = len(eval_runs)
    batch.failed = failed
    batch.average_score = summary.average_score
    batch.keyword_hit_rate = summary.keyword_hit_rate
    batch.source_hit_rate = summary.source_hit_rate
    batch.average_latency_ms = summary.average_latency_ms

    session.add(batch)
    session.commit()
    session.refresh(batch)

    return RagEvalBatchRunResult(
        total_cases=len(eval_cases),
        ran=len(eval_runs),
        failed=failed,
        top_k=request.top_k,
        semantic_weight=semantic_weight,
        keyword_weight=keyword_weight,
        average_score=summary.average_score,
        keyword_hit_rate=summary.keyword_hit_rate,
        source_hit_rate=summary.source_hit_rate,
        average_latency_ms=summary.average_latency_ms,
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-batches",
    response_model=RagEvalBatchesPublic,
    operation_id="read_rag_eval_batches",
)
def read_rag_eval_batches(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 20,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    count_statement = (
        select(func.count())
        .select_from(RagEvalBatch)
        .where(RagEvalBatch.knowledge_base_id == knowledge_base_id)
    )

    count = session.exec(count_statement).one()

    statement = (
        select(RagEvalBatch)
        .where(RagEvalBatch.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalBatch.created_at).desc())
        .offset(skip)
        .limit(limit)
    )

    batches = session.exec(statement).all()

    return RagEvalBatchesPublic(
        data=[rag_eval_batch_to_public(batch) for batch in batches],
        count=count,
    )


@router.delete(
    "/rag-eval-batches/{batch_id}",
    response_model=Message,
    operation_id="delete_rag_eval_batch",
)
def delete_rag_eval_batch(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    batch_id: uuid.UUID,
) -> Any:
    batch = session.get(RagEvalBatch, batch_id)

    if not batch:
        raise HTTPException(status_code=404, detail="RAG eval batch not found")

    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=batch.knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = select(RagEvalRun).where(RagEvalRun.batch_id == batch_id)
    eval_runs = session.exec(statement).all()

    for eval_run in eval_runs:
        eval_run.batch_id = None
        session.add(eval_run)

    session.delete(batch)
    session.commit()

    return Message(message="RAG eval batch deleted successfully")


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-summary",
    response_model=RagEvalSummary,
    operation_id="read_rag_eval_summary",
)
def read_rag_eval_summary(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    limit: int = 200,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = (
        select(RagEvalRun)
        .where(RagEvalRun.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalRun.created_at).desc())
        .limit(limit)
    )

    eval_runs = session.exec(statement).all()

    return build_eval_summary_from_runs(eval_runs)

@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-runs",
    response_model=RagEvalRunsPublic,
    operation_id="read_rag_eval_runs",
)
def read_rag_eval_runs(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    skip: int = 0,
    limit: int = 20,
    keyword: str | None = None,
    failed_only: bool = False,
    keyword_hit: bool | None = None,
    source_hit: bool | None = None,
    error_only: bool = False,
    max_score: float | None = None,
    min_score: float | None = None,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    filters = [
        RagEvalRun.knowledge_base_id == knowledge_base_id,
    ]

    if keyword and keyword.strip():
        pattern = f"%{keyword.strip()}%"
        filters.append(
            or_(
                col(RagEvalRun.question).ilike(pattern),
                col(RagEvalRun.answer).ilike(pattern),
                col(RagEvalRun.expected_source_filename).ilike(pattern),
                col(RagEvalRun.retrieval_type).ilike(pattern),
            )
        )

    if failed_only:
        filters.append(
            or_(
                RagEvalRun.score < 1,
                RagEvalRun.keyword_hit == False,
                RagEvalRun.source_hit == False,
                col(RagEvalRun.error_message).is_not(None),
            )
        )

    if keyword_hit is not None:
        filters.append(RagEvalRun.keyword_hit == keyword_hit)

    if source_hit is not None:
        filters.append(RagEvalRun.source_hit == source_hit)

    if error_only:
        filters.append(col(RagEvalRun.error_message).is_not(None))

    if max_score is not None:
        filters.append(RagEvalRun.score <= max_score)

    if min_score is not None:
        filters.append(RagEvalRun.score >= min_score)

    count_statement = (
        select(func.count())
        .select_from(RagEvalRun)
        .where(*filters)
    )
    count = session.exec(count_statement).one()

    statement = (
        select(RagEvalRun)
        .where(*filters)
        .order_by(col(RagEvalRun.created_at).desc())
        .offset(skip)
        .limit(limit)
    )

    eval_runs = session.exec(statement).all()

    return RagEvalRunsPublic(
        data=[rag_eval_run_to_public(eval_run) for eval_run in eval_runs],
        count=count,
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-failure-analysis",
    response_model=RagEvalFailureAnalysis,
    operation_id="read_rag_eval_failure_analysis",
)
def read_rag_eval_failure_analysis(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    limit: int = 200,
    recent_limit: int = 10,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = (
        select(RagEvalRun)
        .where(RagEvalRun.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalRun.created_at).desc())
        .limit(limit)
    )

    eval_runs = session.exec(statement).all()

    failed_runs = [
        eval_run
        for eval_run in eval_runs
        if is_eval_run_failed(eval_run)
    ]

    total_runs = len(eval_runs)
    failed_count = len(failed_runs)

    keyword_miss_count = sum(
        1 for eval_run in eval_runs if eval_run.keyword_hit is False
    )
    source_miss_count = sum(
        1 for eval_run in eval_runs if eval_run.source_hit is False
    )
    error_count = sum(
        1 for eval_run in eval_runs if eval_run.error_message
    )
    low_score_count = sum(
        1
        for eval_run in eval_runs
        if eval_run.score is not None and eval_run.score < 1
    )
    zero_score_count = sum(
        1
        for eval_run in eval_runs
        if eval_run.score == 0
    )

    average_failed_latency_ms = calculate_average(
        [eval_run.latency_ms for eval_run in failed_runs]
    )

    failure_rate = None

    if total_runs > 0:
        failure_rate = failed_count / total_runs

    recent_failed_runs = failed_runs[:recent_limit]

    return RagEvalFailureAnalysis(
        total_runs=total_runs,
        failed_runs=failed_count,
        failure_rate=failure_rate,
        keyword_miss_count=keyword_miss_count,
        source_miss_count=source_miss_count,
        error_count=error_count,
        low_score_count=low_score_count,
        zero_score_count=zero_score_count,
        average_failed_latency_ms=average_failed_latency_ms,
        recent_failed_runs=[
            rag_eval_run_to_public(eval_run)
            for eval_run in recent_failed_runs
        ],
    )


@router.get(
    "/knowledge-bases/{knowledge_base_id}/rag-eval-param-groups",
    response_model=RagEvalParamGroups,
    operation_id="read_rag_eval_param_groups",
)
def read_rag_eval_param_groups(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    knowledge_base_id: uuid.UUID,
    limit: int = 500,
) -> Any:
    knowledge_base = get_knowledge_base_or_404(
        session=session,
        knowledge_base_id=knowledge_base_id,
    )
    check_knowledge_base_permission(
        knowledge_base=knowledge_base,
        current_user=current_user,
    )

    statement = (
        select(RagEvalRun)
        .where(RagEvalRun.knowledge_base_id == knowledge_base_id)
        .order_by(col(RagEvalRun.created_at).desc())
        .limit(limit)
    )

    eval_runs = session.exec(statement).all()

    preset_statement = (
        select(RagRetrievalPreset)
            .where(RagRetrievalPreset.knowledge_base_id == knowledge_base_id)
    )

    presets = session.exec(preset_statement).all()

    groups = build_eval_param_groups_from_runs(
        eval_runs=eval_runs,
        presets=presets,
    )

    return RagEvalParamGroups(
        data=groups,
        count=len(groups),
    )


@router.get(
    "/documents/{document_id}",
    response_model=DocumentPublic,
)
def read_document(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    document_id: uuid.UUID,
) -> Any:
    document = get_document_or_404(
        session=session,
        document_id=document_id,
    )
    check_document_permission(
        document=document,
        current_user=current_user,
    )

    return document


@router.delete("/documents/{document_id}")
def delete_document(
    *,
    session: SessionDep,
    current_user: CurrentUser,
    document_id: uuid.UUID,
) -> Message:
    document = get_document_or_404(
        session=session,
        document_id=document_id,
    )
    check_document_permission(
        document=document,
        current_user=current_user,
    )

    storage_path = Path(document.storage_path)

    if storage_path.exists() and storage_path.is_file():
        try:
            storage_path.unlink()
        except OSError:
            pass

    session.delete(document)
    session.commit()

    return Message(message="Document deleted successfully")