from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
}


class DocumentParseError(ValueError):
    pass


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower().strip()


def decode_text_file(file_bytes: bytes) -> str:
    encodings = [
        "utf-8",
        "utf-8-sig",
        "gb18030",
        "latin-1",
    ]

    for encoding in encodings:
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise DocumentParseError("无法解析文本文件编码。")


def parse_pdf_file(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception as error:
        raise DocumentParseError(f"PDF 文件读取失败：{error}") from error

    page_texts: list[str] = []

    for page_index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        text = text.strip()

        if text:
            page_texts.append(f"[Page {page_index}]\n{text}")

    parsed_text = "\n\n".join(page_texts).strip()

    if not parsed_text:
        raise DocumentParseError(
            "PDF 中没有提取到可用文本。该 PDF 可能是扫描版图片，需要 OCR。"
        )

    return parsed_text


def parse_docx_file(file_bytes: bytes) -> str:
    try:
        document = DocxDocument(BytesIO(file_bytes))
    except Exception as error:
        raise DocumentParseError(f"Word 文件读取失败：{error}") from error

    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"[Table {table_index}]")

        for row in table.rows:
            row_texts = []

            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)

            if row_texts:
                parts.append(" | ".join(row_texts))

    parsed_text = "\n\n".join(parts).strip()

    if not parsed_text:
        raise DocumentParseError("Word 文件中没有提取到可用文本。")

    return parsed_text


def parse_document_content(
    *,
    filename: str,
    file_bytes: bytes,
) -> str:
    extension = get_file_extension(filename)

    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise DocumentParseError(
            f"暂不支持 {extension or '未知'} 类型文件。"
            "当前支持：txt、md、pdf、docx。"
        )

    if extension in [".txt", ".md"]:
        text = decode_text_file(file_bytes)
    elif extension == ".pdf":
        text = parse_pdf_file(file_bytes)
    elif extension == ".docx":
        text = parse_docx_file(file_bytes)
    else:
        raise DocumentParseError(
            f"暂不支持 {extension} 类型文件。"
        )

    text = text.strip()

    if not text:
        raise DocumentParseError("文件解析后内容为空。")

    return text